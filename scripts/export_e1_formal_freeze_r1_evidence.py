#!/usr/bin/env python3
"""Export E1-FORMAL-FREEZE-R1 evidence package. Does not run formal E1."""
from __future__ import annotations

import json
import shutil
import subprocess
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document

from raven_mcs.utils.hashing import sha256_file
from raven_mcs.utils.serialization import load_json, load_yaml

AUTHORIZED = "53e277c53b01695330652b8e1bc8a234909d56e5"


def _git(root: Path, *args: str) -> str:
    return subprocess.run(
        ["git", *args], cwd=root, check=True,
        capture_output=True, text=True, encoding="utf-8", errors="replace",
    ).stdout


def main() -> int:
    root = Path(__file__).resolve().parents[1]
    deliverables = root / "deliverables"
    deliverables.mkdir(exist_ok=True)
    submission = deliverables / "E1_FORMAL_FREEZE_R1_SUBMISSION"
    if submission.exists():
        shutil.rmtree(submission)
    submission.mkdir(parents=True)

    protocol = load_yaml(root / "configs/frozen/e1_sensorscope_balanced.yaml")
    manifest = load_json(root / "configs/frozen/FROZEN_CONFIG_MANIFEST.json")
    verify = load_json(
        root / "outputs/audits/e1_formal_frozen_manifest_verification.json"
    )
    diff = load_json(
        root / "outputs/audits/e1_formal_execution_diff_scope.json"
    )
    baseline = load_json(
        root / "outputs/validation/e1_formal_freeze_r1_baseline_report.json"
    )
    safety = load_json(
        root / "outputs/validation/e1_formal_freeze_r1_safety_summary.json"
    )
    preflight = load_json(
        root / "outputs/preflight/E1_FORMAL_FREEZE_R1_PREFLIGHT.json"
    )
    resolved = load_json(
        root / "outputs/preflight/e1_formal_resolved_config_hash.json"
    )
    gates = load_json(
        root / "outputs/audits/e1_formal_freeze_r1_gates.json"
    )
    tests = load_json(root / "logs/E1_FORMAL_FREEZE_R1_TEST_SUMMARY.json")
    head = _git(root, "rev-parse", "HEAD").strip()
    clean = _git(root, "status", "--porcelain").strip() == ""
    baseline_path = root / "configs/frozen/e1_selected_baseline.yaml"
    baseline_hash = sha256_file(baseline_path)

    sections = [
        ("Executive Summary", (
            "E1-FORMAL-FREEZE-R1 re-freezes protocol/identity/preflight "
            "evidence only. Formal 25-run E1 has NOT been executed."
        )),
        ("Prior Formal Blocker", (
            "E1-FORMAL-R1 was correctly BLOCKED by freeze/identity mismatches; "
            "historical blocked evidence is retained."
        )),
        ("Authorized Algorithm Identity", AUTHORIZED),
        ("Protocol Parent Identity", AUTHORIZED),
        ("Execution Commit Policy", "runtime_clean_head"),
        ("Frozen local_steps", str(protocol.get("local_steps"))),
        ("Rebuilt Frozen Manifest",
         f"protocol_config_hash={manifest.get('protocol_config_hash')}"),
        ("Selected Baseline Hash Verification",
         f"method={baseline.get('selected_baseline')} "
         f"actual={baseline_hash} "
         f"manifest={manifest['selected_baseline']['file_sha256']} "
         f"match={verify['selected_baseline']['match']}"),
        ("Config Hash Semantics",
         "config_hash = resolved_run_config_hash; "
         "target_group_hash kept separate"),
        ("Execution Diff Scope",
         f"forbidden_core_changes={diff.get('forbidden_core_changes')}"),
        ("Refrozen Protocol Hash",
         str(manifest.get("protocol_config_hash"))),
        ("Resolved Config Hash",
         str(resolved.get("resolved_run_config_hash"))),
        ("Validation Baseline Confirmation",
         json.dumps({
             "selected": baseline.get("selected_baseline"),
             "test_read_count": baseline.get("test_read_count"),
             "local_steps": baseline.get("local_steps"),
         }, sort_keys=True)),
        ("Five-Seed Validation Safety",
         json.dumps({
             "all_seeds_pass": safety.get("all_seeds_pass"),
             "max_first_stage_clip_rate": safety.get(
                 "max_first_stage_clip_rate"
             ),
             "max_second_stage_clip_rate": safety.get(
                 "max_second_stage_clip_rate"
             ),
             "min_median_n_eff": safety.get("min_median_n_eff"),
         }, sort_keys=True)),
        ("Formal Preflight Result",
         f"{preflight.get('status')} mode={preflight.get('mode')}"),
        ("Test and Git Evidence",
         f"tests_ok={tests.get('all_exit_zero')} clean={clean} "
         f"execution_candidate_commit={head}"),
        ("FREEZE-G1 to FREEZE-G10",
         json.dumps(gates.get("gates"), sort_keys=True)),
        ("Remaining Issues",
         "Await teacher authorization of execution_candidate_commit "
         "before starting 25 formal runs."),
        ("Execution Candidate Commit", head),
        ("Teacher Authorization Request",
         "Please authorize clean HEAD as formal E1 execution commit. "
         "Do not rewrite it into the frozen protocol."),
    ]

    md_lines = [
        "# E1-FORMAL-FREEZE-R1 REPORT",
        "",
        f"Generated: {datetime.now(timezone.utc).isoformat()}",
        "",
        "formal 25-run E1 has NOT been executed",
        "E2-E9 have NOT started",
        "",
    ]
    for title, body in sections:
        md_lines.extend([f"## {title}", "", str(body), ""])
    md_path = deliverables / "E1_FORMAL_FREEZE_R1_REPORT.md"
    md_path.write_text("\n".join(md_lines), encoding="utf-8")

    doc = Document()
    doc.add_heading("E1-FORMAL-FREEZE-R1 REPORT", level=1)
    doc.add_paragraph(
        "formal 25-run E1 has NOT been executed; E2-E9 have NOT started"
    )
    for title, body in sections:
        doc.add_heading(title, level=2)
        doc.add_paragraph(str(body))
    docx_path = deliverables / "E1_FORMAL_FREEZE_R1_REPORT.docx"
    doc.save(docx_path)

    readme = deliverables / "E1_FORMAL_FREEZE_R1_SUBMISSION_README.txt"
    readme.write_text(
        "\n".join([
            "RAVEN-MCS E1-FORMAL-FREEZE-R1 submission",
            f"authorized_algorithm_commit={AUTHORIZED}",
            f"protocol_parent_commit={AUTHORIZED}",
            f"execution_candidate_commit={head}",
            f"git_clean={clean}",
            f"local_steps={protocol.get('local_steps')}",
            f"selected_baseline={baseline.get('selected_baseline')}",
            f"protocol_config_hash={manifest.get('protocol_config_hash')}",
            f"resolved_run_config_hash="
            f"{resolved.get('resolved_run_config_hash')}",
            f"preflight={preflight.get('status')}",
            "formal_runs_completed=0",
            "E1=READY_FOR_FINAL_EXECUTION_AUTHORIZATION"
            if gates.get("all_pass") else "E1=NOT EXECUTED",
            "E2-E9=NOT_STARTED",
            "Historical BLOCKED formal report retained under "
            "docs/reports/E1_FORMAL_R1_PREFLIGHT.md",
            "",
        ]),
        encoding="utf-8",
    )

    copy_paths = [
        "configs/frozen/e1_sensorscope_balanced.yaml",
        "configs/frozen/FROZEN_CONFIG_MANIFEST.json",
        "configs/frozen/e1_selected_baseline.yaml",
        "configs/frozen/e1_sensorscope_groups.yaml",
        "configs/frozen/e1_sensorscope_clients.yaml",
        "configs/frozen/e1_pi_target_client_stratum.parquet",
        "configs/frozen/e1_pi_target_manifest.json",
        "configs/frozen/e1_weight_safety.yaml",
        "configs/frozen/e1_protocol_schema.yaml",
        "outputs/audits/e1_formal_frozen_manifest_verification.json",
        "outputs/audits/e1_formal_execution_diff_scope.json",
        "outputs/audits/e1_formal_freeze_r1_gates.json",
        "outputs/preflight/e1_formal_resolved_config.yaml",
        "outputs/preflight/e1_formal_resolved_config_hash.json",
        "outputs/preflight/E1_FORMAL_FREEZE_R1_PREFLIGHT.json",
        "docs/reports/E1_FORMAL_FREEZE_R1_PREFLIGHT.md",
        "docs/reports/E1_FORMAL_FREEZE_R1_BASELINE_AUDIT.md",
        "docs/reports/E1_FORMAL_R1_PREFLIGHT.md",
        "outputs/validation/e1_formal_freeze_r1_baseline_report.json",
        "outputs/validation/e1_formal_freeze_r1_safety_summary.json",
        "logs/E1_FORMAL_FREEZE_R1_TEST_SUMMARY.json",
        "logs/E1_FORMAL_FREEZE_R1_EXACT_COMMANDS.jsonl",
        "STATUS.md",
        "ISSUES.md",
        "CHANGELOG.md",
    ]
    for relative in copy_paths:
        src = root / relative
        if not src.exists():
            continue
        dest = submission / relative
        dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, dest)

    (submission / "git_rev_parse_HEAD.txt").write_text(head + "\n", encoding="utf-8")
    (submission / "git_status_porcelain.txt").write_text(
        _git(root, "status", "--porcelain"), encoding="utf-8",
    )
    (submission / "git_log.txt").write_text(
        _git(root, "log", "-5", "--oneline"), encoding="utf-8",
    )
    shutil.copy2(md_path, submission / md_path.name)
    shutil.copy2(docx_path, submission / docx_path.name)
    shutil.copy2(readme, submission / readme.name)

    zip_path = deliverables / "RAVEN_MCS_E1_FORMAL_FREEZE_R1_EVIDENCE.zip"
    if zip_path.exists():
        zip_path.unlink()
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for path in submission.rglob("*"):
            if path.is_file():
                zf.write(path, path.relative_to(submission).as_posix())

    print(json.dumps({
        "zip": str(zip_path),
        "report": str(docx_path),
        "readme": str(readme),
        "execution_candidate_commit": head,
        "git_clean": clean,
        "zip_sha256": sha256_file(zip_path),
    }, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
