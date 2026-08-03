#!/usr/bin/env python3
"""Export E1-R2 formal freeze-seal evidence without running formal seeds."""
from __future__ import annotations

import argparse
import hashlib
import json
import zipfile
from datetime import datetime, timezone
from html import escape
from pathlib import Path
from typing import Any, Iterable, Mapping

try:
    from docx import Document
except ImportError:  # pragma: no cover - supports minimal export environments
    Document = None  # type: ignore[assignment,misc]

ROOT = Path(__file__).resolve().parents[1]
PACKAGE = "E1_R2_FORMAL_FREEZE_SEAL_R1"
REPORT_NAME = f"{PACKAGE}_REPORT.docx"
README_NAME = f"{PACKAGE}_SUBMISSION_README.txt"
ZIP_NAME = f"RAVEN_MCS_{PACKAGE}_EVIDENCE.zip"
HASH_JSON_NAME = "FINAL_DELIVERABLE_HASHES.json"
HASH_TEXT_NAME = "FINAL_DELIVERABLE_HASHES.txt"

# A requirement is satisfied when at least one of its patterns resolves. Source
# snapshot/bundle/diff are supplemental and are included whenever available.
REQUIRED_EVIDENCE: Mapping[str, tuple[str, ...]] = {
    "source_identity": (
        "**/*E1_R2*SOURCE_IDENTITY*",
        "**/*e1_r2*source_identity*",
    ),
    "source_equivalence": (
        "**/*E1_R2*SOURCE_EQUIVALENCE*",
        "**/*e1_r2*source_equivalence*",
    ),
    "provenance": (
        "**/*E1_R2*PROVENANCE*",
        "**/*e1_r2*provenance*",
    ),
    "pre_export_hashes": (
        "**/*E1_R2*PRE*HASH*",
        "**/*e1_r2*pre*hash*",
    ),
    "post_export_hashes": (
        "**/*E1_R2*POST*HASH*",
        "**/*e1_r2*post*hash*",
    ),
    "frozen_protocol": ("configs/frozen/e1_r2_protocol.yaml",),
    "frozen_manifest": ("configs/frozen/E1_R2_FROZEN_CONFIG_MANIFEST.json",),
    "formal_runner_source": (
        "scripts/run_e1_r2_formal.py",
        "scripts/run_e1_formal.py",
    ),
    "gate_source": (
        "scripts/check_e1_r2_formal_freeze_seal_gates.py",
        "scripts/check_e1_r2_protocol_gates.py",
        "scripts/check_e1_formal_gates.py",
        "scripts/check_e1_run_gates.py",
    ),
    "smoke": (
        "**/*E1_R2*SMOKE*.json",
        "**/*E1_R2*SMOKE*.log",
    ),
    "aggregate_rejection": (
        "**/*E1_R2*AGGREGATE*REJECTION*.json",
        "**/*E1_R2*AGGREGATE*REJECTION*.csv",
    ),
    "junit": (
        "**/*E1_R2*JUNIT*.xml",
        "**/junit*e1_r2*.xml",
    ),
    "logs": (
        "logs/*E1_R2*FORMAL*FREEZE*",
        "logs/*E1_R2*R2FS*",
    ),
    "preflight": (
        "**/*E1_R2*PREFLIGHT*.json",
        "**/*e1_r2*preflight*.json",
    ),
    "noninspection": (
        "outputs/audits/E1_R2_FORMAL_SEED_NONINSPECTION.json",
    ),
    "traces": (
        "data/frozen/e1_r2/formal/ROLE_MANIFEST.json",
        "data/frozen/e1_r2/formal/*/event_trace_manifest.json",
        "data/frozen/e1_r2/formal/**/*",
    ),
    "r2fs_gates": (
        "**/*R2FS*GATE*.json",
        "**/*FORMAL*FREEZE*SEAL*GATE*.json",
    ),
}

OPTIONAL_EVIDENCE: Mapping[str, tuple[str, ...]] = {
    "source_snapshot": (
        "**/*E1_R2*SOURCE_SNAPSHOT*",
        "**/*e1_r2*source_snapshot*",
    ),
    "source_bundle": (
        "**/*E1_R2*SOURCE_BUNDLE*",
        "**/*e1_r2*source_bundle*",
    ),
    "source_diff": (
        "**/*E1_R2*SOURCE_DIFF*",
        "**/*e1_r2*source_diff*",
        "**/*E1_R2*.diff",
        "**/*E1_R2*.patch",
    ),
}

REPORT_SECTIONS = (
    "Executive Summary",
    "Submission Status",
    "Execution Boundary",
    "Source Snapshot and Bundle",
    "Source Diff",
    "Source Identity",
    "Source Equivalence",
    "Provenance",
    "Pre-Export Hashes",
    "Post-Export Hashes",
    "Frozen Protocol",
    "Frozen Manifest",
    "Formal Runner Source",
    "Gate Source",
    "Smoke Test",
    "Aggregate Rejection",
    "JUnit Evidence",
    "Execution Logs",
    "Preflight",
    "Formal-Seed Noninspection",
    "Event Traces",
    "R2FS Gates",
    "Formal Run Status",
)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _read_json(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError):
        return {}
    return value if isinstance(value, dict) else {}


def _write_json(path: Path, value: Mapping[str, Any]) -> None:
    path.write_text(
        json.dumps(dict(value), indent=2, sort_keys=True, default=str) + "\n",
        encoding="utf-8",
    )


def _is_export_output(path: Path, deliverables: Path) -> bool:
    try:
        relative = path.resolve().relative_to(deliverables.resolve())
    except ValueError:
        return False
    return relative.parts[0] in {
        REPORT_NAME,
        README_NAME,
        ZIP_NAME,
        HASH_JSON_NAME,
        HASH_TEXT_NAME,
    }


def _resolve_patterns(
    root: Path,
    patterns: Iterable[str],
    deliverables: Path,
) -> list[Path]:
    found: dict[str, Path] = {}
    for pattern in patterns:
        for path in root.glob(pattern):
            if not path.is_file() or _is_export_output(path, deliverables):
                continue
            relative = path.relative_to(root).as_posix()
            if "/__pycache__/" in f"/{relative}/" or relative.startswith(".git/"):
                continue
            found[relative] = path
    return [found[key] for key in sorted(found)]


def _gate_status(paths: Iterable[Path]) -> str:
    statuses: list[str] = []
    for path in paths:
        value = _read_json(path)
        status = value.get("status")
        if isinstance(status, str):
            statuses.append(status.upper())
        elif value.get("all_pass") is True:
            statuses.append("PASS")
    return "PASS" if "PASS" in statuses else "BLOCKED"


def _section_text(
    title: str,
    *,
    status: str,
    gate_status: str,
    evidence: Mapping[str, list[Path]],
    missing: list[str],
    root: Path,
) -> str:
    def files(*categories: str) -> str:
        paths = [
            path.relative_to(root).as_posix()
            for category in categories
            for path in evidence.get(category, [])
        ]
        return ", ".join(paths) if paths else "Not available."

    texts = {
        "Executive Summary": (
            f"E1-R2 formal freeze-seal evidence status={status}; "
            "this exporter performed no training or formal experiment."
        ),
        "Submission Status": (
            f"status={status}; missing prerequisite categories={len(missing)}"
        ),
        "Execution Boundary": (
            "Packaging only. No runner is imported or invoked. "
            "formal_runs_completed=0."
        ),
        "Source Snapshot and Bundle": files("source_snapshot", "source_bundle"),
        "Source Diff": files("source_diff"),
        "Source Identity": files("source_identity"),
        "Source Equivalence": files("source_equivalence"),
        "Provenance": files("provenance"),
        "Pre-Export Hashes": files("pre_export_hashes"),
        "Post-Export Hashes": files("post_export_hashes"),
        "Frozen Protocol": files("frozen_protocol"),
        "Frozen Manifest": files("frozen_manifest"),
        "Formal Runner Source": files("formal_runner_source"),
        "Gate Source": files("gate_source"),
        "Smoke Test": files("smoke"),
        "Aggregate Rejection": files("aggregate_rejection"),
        "JUnit Evidence": files("junit"),
        "Execution Logs": files("logs"),
        "Preflight": files("preflight"),
        "Formal-Seed Noninspection": files("noninspection"),
        "Event Traces": files("traces"),
        "R2FS Gates": f"status={gate_status}; files={files('r2fs_gates')}",
        "Formal Run Status": "E1-R2 formal runs completed=0; E2-E9=NOT_STARTED.",
    }
    return texts[title]


def _write_minimal_docx(path: Path, sections: list[tuple[str, str]]) -> None:
    paragraphs = [("Title", "E1-R2 Formal Freeze-Seal R1")]
    paragraphs.extend(sections)
    body = "".join(
        f'<w:p><w:pPr><w:pStyle w:val="{style}"/></w:pPr>'
        f'<w:r><w:t xml:space="preserve">{escape(text)}</w:t></w:r></w:p>'
        for style, text in paragraphs
    )
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-'
        'package.relationships+xml"/><Default Extension="xml" '
        'ContentType="application/xml"/><Override PartName="/word/document.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.'
        'document.main+xml"/></Types>'
    )
    relationships = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
        'relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.'
        'org/officeDocument/2006/relationships/officeDocument" '
        'Target="word/document.xml"/></Relationships>'
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/'
        f'2006/main"><w:body>{body}<w:sectPr/></w:body></w:document>'
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", relationships)
        archive.writestr("word/document.xml", document)


def _write_report(path: Path, sections: list[tuple[str, str]]) -> None:
    if Document is None:
        expanded = [
            item
            for title, text in sections
            for item in (("Heading2", title), ("Normal", text))
        ]
        _write_minimal_docx(path, expanded)
        return
    document = Document()
    document.add_heading("E1-R2 Formal Freeze-Seal R1", level=1)
    for title, text in sections:
        document.add_heading(title, level=2)
        document.add_paragraph(text)
    document.save(path)


def export_evidence(root: Path, deliverables: Path) -> dict[str, Any]:
    """Create a complete or graceful partial package from existing evidence."""
    root = root.resolve()
    deliverables = deliverables.resolve()
    deliverables.mkdir(parents=True, exist_ok=True)

    evidence = {
        category: _resolve_patterns(root, patterns, deliverables)
        for category, patterns in {**REQUIRED_EVIDENCE, **OPTIONAL_EVIDENCE}.items()
    }
    missing = [
        category for category in REQUIRED_EVIDENCE if not evidence[category]
    ]
    gate_status = _gate_status(evidence["r2fs_gates"])
    status = "COMPLETE" if not missing and gate_status == "PASS" else "PARTIAL"

    report_path = deliverables / REPORT_NAME
    readme_path = deliverables / README_NAME
    archive_path = deliverables / ZIP_NAME
    hash_json_path = deliverables / HASH_JSON_NAME
    hash_text_path = deliverables / HASH_TEXT_NAME

    sections = [
        (
            title,
            _section_text(
                title,
                status=status,
                gate_status=gate_status,
                evidence=evidence,
                missing=missing,
                root=root,
            ),
        )
        for title in REPORT_SECTIONS
    ]
    _write_report(report_path, sections)
    readme_path.write_text(
        "\n".join(
            [
                "RAVEN-MCS E1-R2-FORMAL-FREEZE-SEAL-R1",
                f"status={status}",
                f"r2fs_gate_status={gate_status}",
                "formal_runs_completed=0",
                "formal_seeds_executed=0",
                "E2_E9=NOT_STARTED",
                f"missing_prerequisite_count={len(missing)}",
                *[f"missing={category}" for category in missing],
                "FINAL_DELIVERABLE_HASHES.json is intentionally outside the ZIP "
                "to avoid a self-hash cycle.",
                "",
            ]
        ),
        encoding="utf-8",
    )

    included: dict[str, Path] = {}
    for paths in evidence.values():
        for path in paths:
            included[path.relative_to(root).as_posix()] = path
    with zipfile.ZipFile(archive_path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.write(report_path, REPORT_NAME)
        archive.write(readme_path, README_NAME)
        for relative, path in sorted(included.items()):
            archive.write(path, relative)

    # The hash manifest is deliberately generated only after the three hashed
    # upload artifacts are final. It is never embedded in the evidence ZIP.
    artifacts = {}
    for path in (archive_path, report_path, readme_path):
        artifacts[path.name] = {
            "sha256": _sha256(path),
            "size_bytes": path.stat().st_size,
        }
    hash_payload = {
        "schema_version": 1,
        "package": PACKAGE,
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "hash_algorithm": "sha256",
        "formal_runs_completed": 0,
        "artifacts": artifacts,
    }
    _write_json(hash_json_path, hash_payload)
    hash_text_path.write_text(
        "".join(
            f"{metadata['sha256']}  {name}\n"
            for name, metadata in artifacts.items()
        ),
        encoding="utf-8",
    )

    return {
        "status": status,
        "r2fs_gate_status": gate_status,
        "formal_runs_completed": 0,
        "missing_prerequisites": missing,
        "evidence_files_included": len(included),
        "report": report_path.as_posix(),
        "readme": readme_path.as_posix(),
        "archive": archive_path.as_posix(),
        "hash_json": hash_json_path.as_posix(),
        "hash_text": hash_text_path.as_posix(),
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=ROOT)
    parser.add_argument("--deliverables", type=Path)
    args = parser.parse_args(argv)
    root = args.root.resolve()
    result = export_evidence(
        root,
        (args.deliverables or root / "deliverables").resolve(),
    )
    print(json.dumps(result, indent=2))
    return 0 if result["status"] == "COMPLETE" else 2


if __name__ == "__main__":
    raise SystemExit(main())
