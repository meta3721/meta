#!/usr/bin/env python3
"""Export E1-R2 protocol-calibration evidence without executing experiments."""
from __future__ import annotations

import argparse
import json
import zipfile
from html import escape
from pathlib import Path
from typing import Any, Mapping

try:
    from docx import Document
except ImportError:  # pragma: no cover - exercised only in minimal environments
    Document = None  # type: ignore[assignment,misc]

ROOT = Path(__file__).resolve().parents[1]
PACKAGE = "E1_R2_PROTOCOL_CALIBRATION_R1"
REPORT_NAME = f"{PACKAGE}_REPORT.docx"
README_NAME = f"{PACKAGE}_SUBMISSION_README.txt"
ZIP_NAME = f"RAVEN_MCS_{PACKAGE}_EVIDENCE.zip"
EVIDENCE_PATHS = (
    "archive/e1_r1_weight_safety_failure/E1_R1_FINAL_STATUS.json",
    "archive/e1_r1_weight_safety_failure/ARCHIVE_HASHES.json",
    "configs/frozen/E1_R2_FROZEN_CONFIG_MANIFEST.json",
    "configs/frozen/e1_r2_protocol.yaml",
    "configs/frozen/e1_r2_weight_safety.yaml",
    "configs/frozen/e1_r2_selected_baseline.yaml",
    "configs/frozen/e1_r2_candidate_selection.json",
    "configs/frozen/e1_r2_seed_registry.yaml",
    "configs/frozen/e1_r2_eventtrace_manifest.json",
    "configs/frozen/e1_r2_analysis.yaml",
    "configs/frozen/e1_r2_scheduler.yaml",
    "configs/e1_r2/candidate_registry.yaml",
    "outputs/e1_r2/calibration/calibration_summary.json",
    "outputs/e1_r2/calibration/passed_candidates.json",
    "outputs/e1_r2/calibration/candidate_gate_matrix.csv",
    "outputs/e1_r2/calibration/candidate_seed_metrics.parquet",
    "outputs/e1_r2/validation/validation_summary.json",
    "outputs/e1_r2/validation/baseline_selection.json",
    "outputs/e1_r2/validation/baseline_selection.parquet",
    "outputs/e1_r2/validation/candidate_seed_metrics.parquet",
    "outputs/e1_r2/validation/candidate_gate_matrix.csv",
    "outputs/e1_r2/validation/final_candidate_selection.json",
    "outputs/audits/E1_R2_PROTOCOL_FREEZE_STATUS.json",
    "outputs/audits/E1_R2_FORMAL_SEED_NONINSPECTION.json",
    "outputs/audits/E1_R2_SCHEDULER_DRY_RUN.json",
    "outputs/audits/E1_R2_PROTOCOL_GATE_REPORT.json",
    "docs/reports/E1_R2_TUNABLE_PARAMETER_INVENTORY.md",
    "logs/E1_R2_PROTOCOL_CALIBRATION_R1_EXACT_COMMANDS.jsonl",
    "STATUS.md",
    "ISSUES.md",
    "CHANGELOG.md",
)
EVIDENCE_DIRS = (
    "data/frozen/e1_r2/calibration",
    "data/frozen/e1_r2/validation",
    "data/frozen/e1_r2/formal",
    "outputs/e1_r2/calibration/runs",
    "outputs/e1_r2/validation/baseline_runs",
    "outputs/e1_r2/validation/candidate_runs",
)


def _read_json(path: Path) -> dict[str, Any]:
    if not path.is_file():
        return {}
    value = json.loads(path.read_text(encoding="utf-8"))
    return value if isinstance(value, dict) else {}


def _write_json(path: Path, value: Mapping[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(dict(value), indent=2, sort_keys=True, default=str) + "\n",
        encoding="utf-8",
    )


def _write_minimal_docx(path: Path, paragraphs: list[str]) -> None:
    """Write a small valid Office Open XML document without python-docx."""
    body = "".join(
        "<w:p><w:r><w:t xml:space=\"preserve\">"
        f"{escape(paragraph)}"
        "</w:t></w:r></w:p>"
        for paragraph in paragraphs
    )
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-'
        'package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" ContentType="application/vnd.'
        'openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        "</Types>"
    )
    relationships = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
        'relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.'
        'org/officeDocument/2006/relationships/officeDocument" '
        'Target="word/document.xml"/></Relationships>'
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/'
        f'2006/main"><w:body>{body}<w:sectPr/></w:body></w:document>'
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", relationships)
        archive.writestr("word/document.xml", document)


def _write_report(
    path: Path,
    *,
    submission_status: str,
    status: str,
    freeze: Mapping[str, Any],
    manifest: Mapping[str, Any],
    noninspection: Mapping[str, Any],
    gates: Mapping[str, Any],
    protocol: Mapping[str, Any],
    selection: Mapping[str, Any],
    validation: Mapping[str, Any],
) -> None:
    summary = (
        f"{submission_status} evidence package; protocol gate status={status}. "
        "This package contains post-selection calibration, validation, freezing, "
        "noninspection, and scheduler dry-run evidence only. It contains no new "
        "formal experiment result."
    )
    freeze_text = (
        f"freeze_status={freeze.get('status', 'BLOCKED')}; "
        f"manifest_status={manifest.get('status', 'ABSENT')}; "
        f"formal_experiments_run={freeze.get('formal_experiments_run', 0)}"
    )
    audit_text = (
        f"status={noninspection.get('status', 'BLOCKED')}; "
        f"training_records={noninspection.get('formal_seed_training_records', 0)}; "
        f"metric_records={noninspection.get('formal_seed_metric_records', 0)}; "
        f"prediction_records="
        f"{noninspection.get('formal_seed_prediction_records', 0)}"
    )
    boundary = (
        "The freeze, audit, gate, scheduler dry-run, and export tools do not "
        "import or call the formal experiment runner. formal_experiments_run=0."
    )
    if Document is None:
        gate_lines = (
            [f"{name}: {gate_status}" for name, gate_status in sorted(gates.items())]
            if gates
            else ["R2P-G1..G10 are BLOCKED pending prerequisite outputs."]
        )
        _write_minimal_docx(
            path,
            [
                "E1-R2 Protocol Calibration R1",
                "Executive Summary",
                summary,
                "Protocol Freeze",
                freeze_text,
                "Formal-Seed Noninspection",
                audit_text,
                "R2P Gates",
                *gate_lines,
                "Execution Boundary",
                boundary,
            ],
        )
        return
    doc = Document()
    doc.add_heading("E1-R2 Protocol Calibration R1", level=1)
    doc.add_heading("Executive Summary", level=2)
    doc.add_paragraph(summary)
    doc.add_heading("Protocol Freeze", level=2)
    doc.add_paragraph(freeze_text)
    doc.add_heading("Formal-Seed Noninspection", level=2)
    doc.add_paragraph(audit_text)
    doc.add_heading("R2P Gates", level=2)
    if gates:
        for name, gate_status in sorted(gates.items()):
            doc.add_paragraph(f"{name}: {gate_status}", style="List Bullet")
    else:
        doc.add_paragraph("R2P-G1..G10 are BLOCKED pending prerequisite outputs.")
    doc.add_heading("Execution Boundary", level=2)
    doc.add_paragraph(boundary)
    doc.add_heading("R2 Main Clip Definition", level=2)
    doc.add_paragraph(
        "Population=observed records; aggregation=global micro per seed; "
        "event=raw u > a_max + 1e-12; threshold=0.05."
    )
    doc.add_heading("Seed Role Separation", level=2)
    doc.add_paragraph(
        f"Calibration={protocol.get('calibration_seeds')}; "
        f"validation={protocol.get('validation_seeds')}; "
        f"formal={protocol.get('formal_seeds')}."
    )
    doc.add_heading("Final Candidate Selection", level=2)
    doc.add_paragraph(
        f"candidate={selection.get('selected_candidate')}; "
        f"a_max={selection.get('selected_a_max')}; "
        f"opportunity forgetting={selection.get('selected_gamma')} unchanged; "
        f"baseline={selection.get('selected_validation_baseline')}."
    )
    doc.add_heading("Validation No-Harm", level=2)
    selected = validation.get("candidates", {}).get(
        selection.get("selected_candidate"), {}
    )
    doc.add_paragraph(
        f"one-sided 95% upper relative RMSE_mu degradation="
        f"{selected.get('one_sided_relative_RMSE_mu_upper_95')}; "
        "threshold=0.03."
    )
    doc.add_heading("Formal Run Status", level=2)
    doc.add_paragraph(
        "E1-R2 formal runs completed=0/25; authorization_status="
        f"{protocol.get('authorization_status')}; E2-E9=NOT_STARTED."
    )
    doc.save(path)


def export_evidence(root: Path, deliverables: Path) -> dict[str, Any]:
    gate = _read_json(root / "outputs/audits/E1_R2_PROTOCOL_GATE_REPORT.json")
    freeze = _read_json(root / "outputs/audits/E1_R2_PROTOCOL_FREEZE_STATUS.json")
    noninspection = _read_json(
        root / "outputs/audits/E1_R2_FORMAL_SEED_NONINSPECTION.json"
    )
    manifest = _read_json(
        root / "configs/frozen/E1_R2_FROZEN_CONFIG_MANIFEST.json"
    )
    selection = _read_json(
        root / "outputs/e1_r2/validation/final_candidate_selection.json"
    )
    validation = _read_json(
        root / "outputs/e1_r2/validation/validation_summary.json"
    )
    protocol_path = root / "configs/frozen/e1_r2_protocol.yaml"
    protocol = {}
    if protocol_path.is_file():
        import yaml
        protocol = yaml.safe_load(protocol_path.read_text(encoding="utf-8")) or {}
    status = gate.get("status", "BLOCKED")
    if status not in {"PASS", "FAIL", "BLOCKED"}:
        status = "BLOCKED"
    submission_status = "COMPLETE" if status == "PASS" else "PARTIAL"
    deliverables.mkdir(parents=True, exist_ok=True)
    report_path = deliverables / REPORT_NAME
    readme_path = deliverables / README_NAME
    archive_path = deliverables / ZIP_NAME

    gates = gate.get("gates", {})
    _write_report(
        report_path,
        submission_status=submission_status,
        status=status,
        freeze=freeze,
        manifest=manifest,
        noninspection=noninspection,
        gates=gates if isinstance(gates, dict) else {},
        protocol=protocol,
        selection=selection,
        validation=validation,
    )

    missing = [
        relative for relative in EVIDENCE_PATHS if not (root / relative).is_file()
    ]
    readme_path.write_text(
        "\n".join(
            [
                "RAVEN-MCS E1-R2-PROTOCOL-CALIBRATION-R1",
                f"status={submission_status}",
                f"gate_status={status}",
                f"freeze_status={freeze.get('status', 'BLOCKED')}",
                f"noninspection_status={noninspection.get('status', 'BLOCKED')}",
                "formal_experiments_run=0",
                "formal_performance_results_created=0",
                f"selected_candidate={selection.get('selected_candidate')}",
                f"selected_a_max={selection.get('selected_a_max')}",
                f"selected_gamma={selection.get('selected_gamma')}",
                f"selected_baseline={selection.get('selected_validation_baseline')}",
                "E1_R2=READY_FOR_TEACHER_REVIEW",
                "E2_E9=NOT_STARTED",
                f"missing_evidence_count={len(missing)}",
                *[f"missing={path}" for path in missing],
                "",
            ]
        ),
        encoding="utf-8",
    )

    included = [
        path
        for path in (root / relative for relative in EVIDENCE_PATHS)
        if path.is_file()
    ]
    with zipfile.ZipFile(archive_path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.write(report_path, REPORT_NAME)
        archive.write(readme_path, README_NAME)
        for path in included:
            archive.write(path, path.relative_to(root).as_posix())
        for relative in EVIDENCE_DIRS:
            directory = root / relative
            if not directory.is_dir():
                continue
            for path in directory.rglob("*"):
                if path.is_file():
                    archive.write(path, path.relative_to(root).as_posix())

    result = {
        "status": submission_status,
        "gate_status": status,
        "formal_experiments_run": 0,
        "report": report_path.as_posix(),
        "readme": readme_path.as_posix(),
        "archive": archive_path.as_posix(),
        "evidence_files_included": len(included),
        "missing_evidence": missing,
    }
    _write_json(deliverables / f"{PACKAGE}_EXPORT_STATUS.json", result)
    return result


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=ROOT)
    parser.add_argument("--deliverables", type=Path)
    args = parser.parse_args(argv)
    root = args.root.resolve()
    result = export_evidence(
        root,
        (args.deliverables or root / "deliverables").resolve(),
    )
    print(json.dumps(result, indent=2))
    return 0 if result["status"] == "COMPLETE" else 2


if __name__ == "__main__":
    raise SystemExit(main())
