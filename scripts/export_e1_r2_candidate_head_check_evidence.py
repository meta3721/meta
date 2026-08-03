#!/usr/bin/env python3
"""Export E1-R2 candidate-head-check evidence without running formal seeds."""
from __future__ import annotations

import argparse
import hashlib
import json
import zipfile
from datetime import datetime, timezone
from html import escape
from pathlib import Path
from typing import Any, Iterable, Mapping

try:
    from docx import Document
except ImportError:  # pragma: no cover
    Document = None  # type: ignore[assignment,misc]

ROOT = Path(__file__).resolve().parents[1]
PACKAGE = "E1_R2_CANDIDATE_HEAD_CHECK_R1"
REPORT_NAME = f"{PACKAGE}_REPORT.docx"
README_NAME = f"{PACKAGE}_SUBMISSION_README.txt"
ZIP_NAME = f"RAVEN_MCS_{PACKAGE}_EVIDENCE.zip"
HASH_JSON_NAME = "FINAL_DELIVERABLE_HASHES.json"
HASH_TEXT_NAME = "FINAL_DELIVERABLE_HASHES.txt"

REQUIRED_EVIDENCE: Mapping[str, tuple[str, ...]] = {
    "exact_head_preflight": (
        "outputs/preflight/E1_R2_EXACT_HEAD_PREFLIGHT.json",
        "outputs/preflight/E1_R2_EXACT_HEAD_PREFLIGHT.md",
    ),
    "exact_head_smoke": (
        "outputs/smoke/e1_r2_exact_head/**/*",
    ),
    "exact_head_replay": (
        "outputs/replay/e1_r2_exact_head/E1_R2_EXACT_HEAD_REPLAY.json",
    ),
    "aggregate_rejection": (
        "logs/e1_r2_candidate_head/*aggregate*",
        "logs/E1_R2_CANDIDATE_HEAD_CHECK_R1_EXACT_COMMANDS.jsonl",
    ),
    "candidate_head_gates": (
        "outputs/audits/E1_R2_CANDIDATE_HEAD_GATES.json",
        "outputs/audits/E1_R2_CANDIDATE_HEAD_GATES.md",
    ),
    "frozen_protocol": ("configs/frozen/e1_r2_protocol.yaml",),
    "frozen_manifest": ("configs/frozen/E1_R2_FROZEN_CONFIG_MANIFEST.json",),
    "selected_baseline": ("configs/frozen/e1_r2_selected_baseline.yaml",),
    "target_hash_semantics": (
        "docs/reports/E1_TARGET_HASH_SEMANTICS.md",
        "outputs/diagnostics/E1_TARGET_HASH_SEMANTICS.json",
    ),
    "gitattributes": (".gitattributes",),
    "preflight_source": ("scripts/check_e1_r2_formal_preflight.py",),
    "replay_source": ("scripts/replay_e1_r2_calval_source_identity.py",),
    "gate_source": ("scripts/check_e1_r2_candidate_head_gates.py",),
    "export_source": ("scripts/export_e1_r2_candidate_head_check_evidence.py",),
    "logged_runner_source": ("scripts/run_e1_r2_candidate_head_logged.py",),
    "command_ledger": (
        "logs/E1_R2_CANDIDATE_HEAD_CHECK_R1_EXACT_COMMANDS.jsonl",
    ),
    "bundle_identity": (
        "outputs/audits/E1_R2_CANDIDATE_HEAD_BUNDLE_IDENTITY.json",
        "outputs/audits/e1_r2_source_identity/E1_R2_C1_SOURCE_IDENTITY.json",
    ),
    "unit_tests": (
        "tests/unit/test_e1_r2_candidate_head_identity.py",
    ),
}

OPTIONAL_EVIDENCE: Mapping[str, tuple[str, ...]] = {
    "calibration_audit": (
        "outputs/audits/e1_r3_calibration_summary.json",
        "configs/frozen/e1_measurement_calibration_manifest.json",
    ),
    "renormalize_helper": ("scripts/renormalize_text_eol_lf.py",),
}

REPORT_SECTIONS = (
    "Executive Summary",
    "Submission Status",
    "Execution Boundary",
    "Candidate Commit Binding",
    "Exact-Head Preflight",
    "Protocol File Hash",
    "Protocol Payload Hash",
    "Target Hash Semantics",
    "Selected Candidate Identity",
    "Baseline Path and Hash",
    "Formal Seeds",
    "Clip Fields",
    "Exact-Head Smoke",
    "Aggregate Rejection",
    "Exact-Head Replay",
    "Candidate Head Gates",
    "Bundle Identity",
    "Git Clean HEAD Match",
    "Command Ledger",
    "Unit Test Coverage",
    "LF Checkout Policy",
    "Formal Run Status",
)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _read_json(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError):
        return {}
    return value if isinstance(value, dict) else {}


def _write_json(path: Path, value: Mapping[str, Any]) -> None:
    path.write_text(
        json.dumps(dict(value), indent=2, sort_keys=True, default=str) + "\n",
        encoding="utf-8",
    )


def _is_export_output(path: Path, deliverables: Path) -> bool:
    try:
        relative = path.resolve().relative_to(deliverables.resolve())
    except ValueError:
        return False
    return relative.parts[0] in {
        REPORT_NAME,
        README_NAME,
        ZIP_NAME,
        HASH_JSON_NAME,
        HASH_TEXT_NAME,
    }


def _resolve_patterns(
    root: Path,
    patterns: Iterable[str],
    deliverables: Path,
) -> list[Path]:
    found: dict[str, Path] = {}
    for pattern in patterns:
        for path in root.glob(pattern):
            if not path.is_file() or _is_export_output(path, deliverables):
                continue
            relative = path.relative_to(root).as_posix()
            if "/__pycache__/" in f"/{relative}/" or relative.startswith(".git/"):
                continue
            found[relative] = path
    return [found[key] for key in sorted(found)]


def _gate_status(paths: Iterable[Path]) -> str:
    statuses: list[str] = []
    for path in paths:
        value = _read_json(path)
        status = value.get("status")
        if isinstance(status, str):
            statuses.append(status.upper())
        elif value.get("all_pass") is True:
            statuses.append("PASS")
    return "PASS" if "PASS" in statuses else "BLOCKED"


def _section_text(
    title: str,
    *,
    status: str,
    gate_status: str,
    evidence: Mapping[str, list[Path]],
    missing: list[str],
    root: Path,
) -> str:
    def files(*categories: str) -> str:
        paths = [
            path.relative_to(root).as_posix()
            for category in categories
            for path in evidence.get(category, [])
        ]
        return ", ".join(paths) if paths else "Not available."

    preflight = _read_json(root / "outputs/preflight/E1_R2_EXACT_HEAD_PREFLIGHT.json")
    replay = _read_json(
        root / "outputs/replay/e1_r2_exact_head/E1_R2_EXACT_HEAD_REPLAY.json"
    )
    texts = {
        "Executive Summary": (
            f"E1-R2 candidate-head-check evidence status={status}; "
            "this exporter performed no training or formal experiment."
        ),
        "Submission Status": (
            f"status={status}; missing prerequisite categories={len(missing)}; "
            f"head_gate_status={gate_status}"
        ),
        "Execution Boundary": (
            "Packaging only. No formal seeds executed. formal_runs_completed=0."
        ),
        "Candidate Commit Binding": (
            f"expected_commit={preflight.get('expected_commit')}; "
            f"runtime_head={preflight.get('runtime_head')}; "
            f"replay_candidate={replay.get('candidate_commit')}"
        ),
        "Exact-Head Preflight": files("exact_head_preflight"),
        "Protocol File Hash": (
            f"protocol_file_hash={preflight.get('protocol_file_hash')}; "
            "definition=checkout bytes SHA-256"
        ),
        "Protocol Payload Hash": (
            f"protocol_payload_hash={preflight.get('protocol_payload_hash')}; "
            "definition=normalized YAML object SHA-256 (primary identity)"
        ),
        "Target Hash Semantics": files("target_hash_semantics"),
        "Selected Candidate Identity": (
            "selected_candidate=C2; selected_a_max=40.0; "
            "selected_opportunity_forgetting=0.95"
        ),
        "Baseline Path and Hash": (
            f"path={preflight.get('selected_baseline_path')}; "
            f"file_hash={preflight.get('selected_baseline_file_hash')}; "
            f"files={files('selected_baseline')}"
        ),
        "Formal Seeds": "formal_seeds=[28001,28002,28003,28004,28005]",
        "Clip Fields": (
            "main_clip_population=observed_records; "
            "main_clip_aggregation=global_micro_per_seed; "
            "main_clip_threshold=0.05"
        ),
        "Exact-Head Smoke": files("exact_head_smoke"),
        "Aggregate Rejection": files("aggregate_rejection"),
        "Exact-Head Replay": (
            f"replay_status={replay.get('replay_status', replay.get('status'))}; "
            "unexplained_numeric_difference_count="
            f"{replay.get('unexplained_numeric_difference_count')}; "
            f"files={files('exact_head_replay')}"
        ),
        "Candidate Head Gates": (
            f"status={gate_status}; files={files('candidate_head_gates')}"
        ),
        "Bundle Identity": files("bundle_identity"),
        "Git Clean HEAD Match": (
            f"runtime_git_clean={preflight.get('runtime_git_clean')}; "
            f"expected_commit_match="
            f"{(preflight.get('checks') or {}).get('expected_commit_match')}"
        ),
        "Command Ledger": files("command_ledger"),
        "Unit Test Coverage": files("unit_tests"),
        "LF Checkout Policy": files("gitattributes", "renormalize_helper"),
        "Formal Run Status": "E1-R2 formal runs completed=0; E2-E9=NOT_STARTED.",
    }
    return texts[title]


def _write_minimal_docx(path: Path, sections: list[tuple[str, str]]) -> None:
    paragraphs = [("Title", "E1-R2 Candidate Head Check R1")]
    paragraphs.extend(sections)
    body = "".join(
        f'<w:p><w:pPr><w:pStyle w:val="{style}"/></w:pPr>'
        f'<w:r><w:t xml:space="preserve">{escape(text)}</w:t></w:r></w:p>'
        for style, text in paragraphs
    )
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-'
        'package.relationships+xml"/><Default Extension="xml" '
        'ContentType="application/xml"/><Override PartName="/word/document.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.'
        'document.main+xml"/></Types>'
    )
    relationships = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
        'relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.'
        'org/officeDocument/2006/relationships/officeDocument" '
        'Target="word/document.xml"/></Relationships>'
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/'
        f'2006/main"><w:body>{body}<w:sectPr/></w:body></w:document>'
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", relationships)
        archive.writestr("word/document.xml", document)


def _write_report(path: Path, sections: list[tuple[str, str]]) -> None:
    if Document is None:
        expanded = [
            item
            for title, text in sections
            for item in (("Heading2", title), ("Normal", text))
        ]
        _write_minimal_docx(path, expanded)
        return
    document = Document()
    document.add_heading("E1-R2 Candidate Head Check R1", level=1)
    for title, text in sections:
        document.add_heading(title, level=2)
        document.add_paragraph(text)
    document.save(path)


def export_evidence(root: Path, deliverables: Path) -> dict[str, Any]:
    root = root.resolve()
    deliverables = deliverables.resolve()
    deliverables.mkdir(parents=True, exist_ok=True)

    evidence = {
        category: _resolve_patterns(root, patterns, deliverables)
        for category, patterns in {**REQUIRED_EVIDENCE, **OPTIONAL_EVIDENCE}.items()
    }
    missing = [
        category for category in REQUIRED_EVIDENCE if not evidence[category]
    ]
    gate_status = _gate_status(evidence["candidate_head_gates"])
    status = "COMPLETE" if not missing and gate_status == "PASS" else "PARTIAL"

    report_path = deliverables / REPORT_NAME
    readme_path = deliverables / README_NAME
    archive_path = deliverables / ZIP_NAME
    hash_json_path = deliverables / HASH_JSON_NAME
    hash_text_path = deliverables / HASH_TEXT_NAME

    sections = [
        (
            title,
            _section_text(
                title,
                status=status,
                gate_status=gate_status,
                evidence=evidence,
                missing=missing,
                root=root,
            ),
        )
        for title in REPORT_SECTIONS
    ]
    assert len(REPORT_SECTIONS) == 22
    _write_report(report_path, sections)
    preflight = _read_json(root / "outputs/preflight/E1_R2_EXACT_HEAD_PREFLIGHT.json")
    readme_path.write_text(
        "\n".join(
            [
                "RAVEN-MCS E1-R2-CANDIDATE-HEAD-CHECK-R1",
                f"status={status}",
                f"head_gate_status={gate_status}",
                f"candidate_commit={preflight.get('expected_commit')}",
                "formal_runs_completed=0",
                "formal_seeds_executed=0",
                "E2_E9=NOT_STARTED",
                f"missing_prerequisite_count={len(missing)}",
                *[f"missing={category}" for category in missing],
                "FINAL_DELIVERABLE_HASHES.json is intentionally outside the ZIP "
                "to avoid a self-hash cycle.",
                "",
            ]
        ),
        encoding="utf-8",
    )

    included: dict[str, Path] = {}
    for paths in evidence.values():
        for path in paths:
            included[path.relative_to(root).as_posix()] = path
    with zipfile.ZipFile(archive_path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.write(report_path, REPORT_NAME)
        archive.write(readme_path, README_NAME)
        for relative, path in sorted(included.items()):
            archive.write(path, relative)

    artifacts = {}
    for path in (archive_path, report_path, readme_path):
        artifacts[path.name] = {
            "sha256": _sha256(path),
            "size_bytes": path.stat().st_size,
        }
    hash_payload = {
        "schema_version": 1,
        "package": PACKAGE,
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "hash_algorithm": "sha256",
        "formal_runs_completed": 0,
        "candidate_commit": preflight.get("expected_commit"),
        "artifacts": artifacts,
    }
    _write_json(hash_json_path, hash_payload)
    hash_text_path.write_text(
        "".join(
            f"{metadata['sha256']}  {name}\n"
            for name, metadata in artifacts.items()
        ),
        encoding="utf-8",
    )

    return {
        "status": status,
        "head_gate_status": gate_status,
        "formal_runs_completed": 0,
        "missing_prerequisites": missing,
        "evidence_files_included": len(included),
        "report_sections": len(REPORT_SECTIONS),
        "report": report_path.as_posix(),
        "readme": readme_path.as_posix(),
        "archive": archive_path.as_posix(),
        "hash_json": hash_json_path.as_posix(),
        "hash_text": hash_text_path.as_posix(),
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=ROOT)
    parser.add_argument("--deliverables", type=Path)
    args = parser.parse_args(argv)
    root = args.root.resolve()
    result = export_evidence(
        root,
        (
            args.deliverables
            or root / "deliverables/TO_SUBMIT_E1_R2_CANDIDATE_HEAD_CHECK_R1"
        ).resolve(),
    )
    print(json.dumps(result, indent=2))
    return 0 if result["status"] == "COMPLETE" else 2


if __name__ == "__main__":
    raise SystemExit(main())
