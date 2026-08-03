#!/usr/bin/env python3
"""Generate the formal Word report for E1-R2 results evidence seal fix."""
from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

import pandas as pd

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt
except ImportError:  # pragma: no cover
    Document = None  # type: ignore[assignment,misc]

ROOT = Path(__file__).resolve().parents[1]
PACKAGE = "E1_R2_RESULTS_EVIDENCE_SEAL_FIX_R1"
FORMAL_COMMIT = "e8bd1fc777431c2609def257a04fba093f0daf24"
PARENT_AUDIT = "38210700bdb6b3365930a1a4ed92a2488913d874"


def _json(path: Path) -> dict[str, Any]:
    if not path.is_file():
        return {}
    value = json.loads(path.read_text(encoding="utf-8"))
    return value if isinstance(value, dict) else {}


def _add_table(document: Any, frame: pd.DataFrame, title: str) -> None:
    document.add_heading(title, level=2)
    if frame.empty:
        document.add_paragraph("Table unavailable.")
        return
    table = document.add_table(rows=1, cols=len(frame.columns))
    table.style = "Table Grid"
    for index, column in enumerate(frame.columns):
        table.rows[0].cells[index].text = str(column)
    for row in frame.itertuples(index=False):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            if isinstance(value, float):
                cells[index].text = f"{value:.6g}"
            else:
                cells[index].text = str(value)
    document.add_paragraph("")


def _add_figure(document: Any, path: Path, caption: str) -> None:
    document.add_heading(caption, level=2)
    if path.is_file():
        document.add_picture(str(path), width=Inches(5.8))
        paragraph = document.add_paragraph(caption)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        document.add_paragraph(f"[Missing figure: {path.name}]")


def build_report(root: Path, output_path: Path | None = None) -> dict[str, Any]:
    root = Path(root).resolve()
    if output_path is None:
        output_path = root / f"deliverables/TO_SUBMIT_{PACKAGE}" / f"{PACKAGE}_REPORT.docx"
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    if Document is None:
        raise RuntimeError("python-docx is required for the formal report")

    stats = root / "outputs/statistics/E1_R2_FINAL_SEALED"
    paper = root / "outputs/paper/E1_R2_FINAL"
    identity = _json(root / "outputs/audits/E1_R2_RESULTS_EVIDENCE_SEAL_IDENTITY.json")
    imm = _json(root / "outputs/audits/E1_R2_RESULTS_EVIDENCE_FIX_IMMUTABILITY_CHECK.json")
    no_harm = _json(stats / "no_harm_summary.json")
    fix_gates = _json(root / "outputs/gates/E1_R2_RESULTS_EVIDENCE_FIX/E1_R2_RESULTS_EVIDENCE_FIX_GATES.json")
    sealed = _json(root / "outputs/gates/E1_R2_SEALED/E1_R2_SEALED_GATES.json")
    recompute = _json(root / "outputs/audits/E1_R2_FORMAL_RESULTS_INDEPENDENT_RECOMPUTE.json")
    solver = _json(root / "outputs/audits/E1_R2_RAVEN_SOLVER_RESIDUAL_MAXIMA.json")
    semantic = _json(root / "outputs/audits/E1_R2_FORMAL_SEMANTIC_AUDIT.json")
    family = _json(stats / "holm_family_registry.json")
    holm = pd.read_csv(stats / "holm_results.csv") if (stats / "holm_results.csv").is_file() else pd.DataFrame()
    main = pd.read_csv(paper / "tables/table_e1_main_metrics.csv") \
        if (paper / "tables/table_e1_main_metrics.csv").is_file() else pd.DataFrame()
    safety = pd.read_csv(paper / "tables/table_e1_safety.csv") \
        if (paper / "tables/table_e1_safety.csv").is_file() else pd.DataFrame()
    stats_table = pd.read_csv(paper / "tables/table_e1_wilcoxon_holm.csv") \
        if (paper / "tables/table_e1_wilcoxon_holm.csv").is_file() else pd.DataFrame()
    fallback = pd.read_csv(paper / "tables/table_e1_solver_fallback.csv") \
        if (paper / "tables/table_e1_solver_fallback.csv").is_file() else pd.DataFrame()
    runtime = pd.read_csv(paper / "tables/table_e1_runtime_updates.csv") \
        if (paper / "tables/table_e1_runtime_updates.csv").is_file() else pd.DataFrame()

    raven_nh = no_harm.get("no_harm_tests", {}).get("raven", {})
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    document.add_heading("RAVEN-MCS E1-R2 Results Evidence Seal Fix R1", level=0)
    cover = document.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.add_run(
        "Formal results evidence package for teacher review\n"
        f"formal_execution_commit = {FORMAL_COMMIT}\n"
        f"parent_results_audit_commit = {PARENT_AUDIT}\n"
        f"results_evidence_seal_commit = {identity.get('results_evidence_seal_commit', 'PENDING')}\n"
        f"Status = {fix_gates.get('e1_r2_final_status', 'IN_PROGRESS')}"
    )

    document.add_heading("1. Executive Summary", level=1)
    document.add_paragraph(
        "This package reseals the completed E1-R2 formal 25-run matrix without retraining. "
        "Holm correction is applied as one family per metric (size 4). "
        "RAVEN ranks second behind TimeAlign, passes the 3% no-harm gate, and does not "
        "establish Holm-corrected statistical superiority. The evidence ZIP is intended to "
        "be self-contained for SEALED-gate replay."
    )

    document.add_heading("2. Formal Experiment Identity", level=1)
    document.add_paragraph(
        "Protocol E1-R2 / candidate C2 / a_max=40 / opportunity_forgetting=0.95 / "
        "baseline=flamf_timealign_adapted / local_steps=2 / windows=100 / clients=8 / S_max=5. "
        f"Formal seeds 28001–28005. Immutable check status={imm.get('status')} with "
        f"hash_mismatch_count={imm.get('hash_mismatch_count')}."
    )

    document.add_heading("3. 25-Run Completion Matrix", level=1)
    document.add_paragraph("Formal runs completed = 25/25; failed = 0; exact restart count = 0.")

    _add_table(document, main, "4. Main Results Table")
    _add_figure(
        document,
        paper / "figures/fig_e1_rmse_mu_by_method.png",
        "Figure 1. Mean RMSE_mu by method (n=5, seed SD error bars).",
    )
    _add_figure(
        document,
        paper / "figures/fig_e1_relative_degradation.png",
        "Figure 2. RAVEN versus TimeAlign relative degradation.",
    )

    document.add_heading("5. No-Harm Recalculation", level=1)
    document.add_paragraph(
        f"Mean relative degradation={raven_nh.get('relative_degradation_mean')}; "
        f"one-sided 95% upper bound={raven_nh.get('one_sided_upper_bound')}; "
        f"threshold=0.03; no_harm_pass={raven_nh.get('no_harm_pass')}."
    )
    document.add_heading("6. Corrected No-Harm Gate Logic", level=1)
    document.add_paragraph(
        "formal_no_harm_conclusion is true only when no_harm_tests.raven.no_harm_pass is true "
        "and the finite one-sided upper bound is strictly below 0.03. Formal mode never "
        "auto-passes missing, NaN, wrong-baseline, or incomplete-seed cases."
    )

    _add_table(document, stats_table, "7. Wilcoxon and Metric-Wise Holm Results")
    document.add_heading("8. Metric-Wise Holm Family Definition", level=1)
    document.add_paragraph(
        f"family_mode={family.get('family_mode')}; n_families={family.get('n_families')}. "
        "Each of RMSE_mu, RMSE_rho, Gap_mis, Tail_RMSE, and runtime forms an independent "
        "family of the four RAVEN-versus-baseline comparisons."
    )
    document.add_heading("9. Statistical Power Limitation", level=1)
    document.add_paragraph(
        "With n=5 paired seeds, non-significant Holm results do not prove equivalence; "
        "they only mean superiority is not established under the pre-registered family."
    )

    _add_table(document, safety, "10. First/Second-Stage Clip Safety and ESS")
    document.add_paragraph(
        f"Semantic totals: q leakage={semantic.get('total_q_nonattempt_leakage')}, "
        f"failed-attempt omission={semantic.get('total_failed_attempt_omission')}, "
        f"unsupported arrival={semantic.get('total_unsupported_arrival_contribution')}, "
        f"solver failure={semantic.get('total_solver_failure')}, "
        f"nan/inf={semantic.get('total_nan_inf')}."
    )
    _add_table(document, fallback, "11. Solver Fallback Table")
    document.add_paragraph(
        f"Total fallback invocations={solver.get('total_fallback_invoked')}; "
        f"complete solver failures={solver.get('total_solver_failures')}."
    )
    _add_table(document, runtime, "12. Runtime and Communication Updates")
    document.add_paragraph(
        "Communication is the count of received client updates, not byte volume."
    )

    document.add_heading("13. Claims Supported", level=1)
    document.add_paragraph(
        "TimeAlign is first on mean RMSE_mu; RAVEN is second; RAVEN passes 3% no-harm; "
        "RAVEN has better mean RMSE_mu than FedAvg/FedAsync/TwoStage; all safety/semantic "
        "gates pass; communication is update count."
    )
    document.add_heading("14. Claims Not Supported", level=1)
    document.add_paragraph(
        "RAVEN is not best on RMSE_mu; Holm-corrected superiority is not established; "
        "communication bytes are not measured; fallback-free solving is false."
    )

    document.add_heading("15. FIX-G1 to FIX-G10", level=1)
    for gate, status in fix_gates.get("gates", {}).items():
        document.add_paragraph(f"{gate}: {status}", style="List Bullet")
    document.add_heading("16. Evidence Self-Containment", level=1)
    document.add_paragraph(
        f"Self-contained replay status={identity.get('self_contained_replay_status')}. "
        "Unpacked evidence ZIP can run scripts/check_e1_r2_results_seal_gates.py --root ."
    )
    document.add_heading("17. Audit Commit Identity", level=1)
    document.add_paragraph(
        f"results_evidence_seal_commit={identity.get('results_evidence_seal_commit')}; "
        f"bundle_verify={identity.get('bundle_verify_status')}; "
        f"git_status_clean={identity.get('git_status_clean')}."
    )
    document.add_heading("18. Final Status", level=1)
    document.add_paragraph(
        f"E1-R2={fix_gates.get('e1_r2_final_status', sealed.get('e1_r2_final_status'))}; "
        "statistical superiority=NOT_ESTABLISHED; E2-E9=NOT_STARTED."
    )
    document.add_heading("19. Next Step", level=1)
    document.add_paragraph(
        "Stop for teacher review. Do not start E2–E9 and do not retrain formal runs."
    )

    # Keep method means readable without dumping the whole recompute JSON.
    document.add_heading("Appendix A. Method Mean RMSE_mu", level=1)
    means = pd.DataFrame(recompute.get("method_summary", []))
    if not means.empty:
        cols = [c for c in ("method", "RMSE_mu_mean", "RMSE_mu_std", "RMSE_mu_rank") if c in means.columns]
        _add_table(document, means[cols], "Appendix Table A1")

    if not holm.empty:
        rmse = holm[holm["metric"] == "RMSE_mu"][
            [c for c in (
                "comparison", "family_id", "family_size", "p_value",
                "holm_adjusted_p", "holm_significant",
            ) if c in holm.columns]
        ]
        _add_table(document, rmse, "Appendix Table A2. RMSE_mu Holm Detail")

    document.save(str(output_path))
    return {"status": "PASS", "report": output_path.as_posix()}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=ROOT)
    parser.add_argument("--output", type=Path, default=None)
    args = parser.parse_args(argv)
    result = build_report(args.root, args.output)
    print(json.dumps(result, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
