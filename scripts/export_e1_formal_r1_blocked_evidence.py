#!/usr/bin/env python3
"""Export E1-FORMAL-R1 blocked preflight report and evidence package.

Does not execute formal runs. Documents freeze/identity blockers only.
"""
from __future__ import annotations

import json
import platform
import shutil
import subprocess
import sys
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document

from raven_mcs.utils.hashing import sha256_file
from raven_mcs.utils.serialization import load_json, load_yaml

REQUIRED_COMMIT = "53e277c53b01695330652b8e1bc8a234909d56e5"
METHODS = (
    "fedavg_window",
    "fedasync_window",
    "flamf_timealign_adapted",
    "twostage_hajek",
    "raven",
)
SEEDS = (26001, 26002, 26003, 26004, 26005)


def _git(root: Path, *args: str) -> str:
    return subprocess.run(
        ["git", *args],
        cwd=root,
        check=True,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    ).stdout


def collect_preflight(root: Path) -> dict:
    head = _git(root, "rev-parse", "HEAD").strip()
    porcelain = _git(root, "status", "--porcelain")
    git_clean = porcelain.strip() == ""

    files = {
        "protocol": root / "configs/frozen/e1_sensorscope_balanced.yaml",
        "groups": root / "configs/frozen/e1_sensorscope_groups.yaml",
        "clients": root / "configs/frozen/e1_sensorscope_clients.yaml",
        "pi_target": root / "configs/frozen/e1_pi_target_client_stratum.parquet",
        "pi_manifest": root / "configs/frozen/e1_pi_target_manifest.json",
        "baseline": root / "configs/frozen/e1_selected_baseline.yaml",
        "weight_safety": root / "configs/frozen/e1_weight_safety.yaml",
        "frozen_manifest": root / "configs/frozen/FROZEN_CONFIG_MANIFEST.json",
        "trace_manifest": root / "outputs/event_traces/E1_BALANCED_EVENTTRACE_MANIFEST.json",
    }
    exists = {key: path.exists() for key, path in files.items()}
    hashes = {
        key: (sha256_file(path) if path.exists() else None)
        for key, path in files.items()
    }
    protocol = load_yaml(files["protocol"]) if exists["protocol"] else {}
    frozen = load_json(files["frozen_manifest"]) if exists["frozen_manifest"] else {}
    baseline = load_yaml(files["baseline"]) if exists["baseline"] else {}
    trace_manifest = (
        load_json(files["trace_manifest"]) if exists["trace_manifest"] else {}
    )
    pi_manifest = load_json(files["pi_manifest"]) if exists["pi_manifest"] else {}

    trace_rows = []
    for seed in SEEDS:
        directory = root / f"outputs/event_traces/e1_balanced_seed{seed}"
        events = directory / "events.parquet"
        audit = directory / "audit.json"
        manifest = directory / "event_trace_manifest.json"
        audit_obj = load_json(audit) if audit.exists() else {}
        manifest_obj = load_json(manifest) if manifest.exists() else {}
        expected = frozen.get("event_trace_hashes", {}).get(str(seed))
        actual = manifest_obj.get("event_trace_hash")
        if not actual and str(seed) in (trace_manifest.get("traces") or {}):
            actual = trace_manifest["traces"][str(seed)]["event_trace_hash"]
        generation = (
            manifest_obj.get("generation_git_commit")
            or (trace_manifest.get("traces") or {})
            .get(str(seed), {})
            .get("generation_git_commit")
        )
        trace_rows.append(
            {
                "seed": seed,
                "dir_exists": directory.exists(),
                "events_exists": events.exists(),
                "events_sha256": sha256_file(events) if events.exists() else None,
                "audit_pass": bool(
                    audit_obj.get(
                        "hard_gate_pass",
                        audit_obj.get(
                            "audit_pass", audit_obj.get("pass", False),
                        ),
                    )
                ),
                "expected_hash": expected,
                "actual_hash": actual,
                "hash_match": bool(expected and actual and expected == actual),
                "generation_git_commit": generation,
            }
        )

    req = root / "requirements.txt"
    lock = root / "requirements.lock"
    dep_path = lock if lock.exists() else req if req.exists() else None
    dep_hash = sha256_file(dep_path) if dep_path else None
    free_gb = shutil.disk_usage(root).free / (1024**3)

    test_summary_path = root / "logs/E1_ENTRY_R4_TEST_SUMMARY.json"
    test_summary = load_json(test_summary_path) if test_summary_path.exists() else {}

    packaging_paths = {
        "STATUS.md",
        "ISSUES.md",
        ".gitignore",
        "docs/reports/E1_FORMAL_R1_PREFLIGHT.md",
        "scripts/export_e1_formal_r1_blocked_evidence.py",
        "deliverables/E1_FORMAL_REPORT.md",
        "deliverables/E1_FORMAL_REPORT.docx",
        "deliverables/E1_FORMAL_SUBMISSION_README.txt",
        "deliverables/E1_FORMAL_DELIVERABLE_HASHES.json",
        "deliverables/RAVEN_MCS_E1_FORMAL_EVIDENCE.zip",
        "outputs/audits/e1_formal_r1_preflight.json",
    }
    dirty_paths = [
        line[3:].replace("\\", "/").strip()
        for line in porcelain.splitlines()
        if line.strip()
    ]
    non_packaging_dirty = [
        path for path in dirty_paths
        if path not in packaging_paths
        and not path.startswith("outputs/evidence_formal_r1/")
        and not path.startswith("deliverables/")
    ]

    blockers: list[str] = []
    if head != REQUIRED_COMMIT:
        blockers.append(f"HEAD {head} != required {REQUIRED_COMMIT}")
    if non_packaging_dirty:
        blockers.append(
            "git worktree has non-packaging dirty paths: "
            + ", ".join(non_packaging_dirty)
        )
    if protocol.get("git_commit") != REQUIRED_COMMIT:
        blockers.append(
            "frozen protocol git_commit="
            f"{protocol.get('git_commit')} != required {REQUIRED_COMMIT}"
        )
    if frozen.get("git_commit") != REQUIRED_COMMIT:
        blockers.append(
            "FROZEN_CONFIG_MANIFEST git_commit="
            f"{frozen.get('git_commit')} != required {REQUIRED_COMMIT}"
        )
    if "local_steps" not in protocol:
        blockers.append("frozen protocol missing local_steps")
    missing = [key for key, value in exists.items() if not value]
    if missing:
        blockers.append("missing frozen/evidence files: " + ", ".join(missing))
    if any(not row["dir_exists"] or not row["events_exists"] for row in trace_rows):
        blockers.append("missing EventTrace artifacts")
    if any(not row["audit_pass"] for row in trace_rows):
        blockers.append("EventTrace audit not all pass")

    return {
        "status": "BLOCKED" if blockers else "PASS",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "required_commit": REQUIRED_COMMIT,
        "head_commit": head,
        "git_clean": git_clean,
        "git_dirty_paths": dirty_paths,
        "git_non_packaging_dirty_paths": non_packaging_dirty,
        "git_clean_except_packaging": len(non_packaging_dirty) == 0,
        "python_version": sys.version.split()[0],
        "venv_python": sys.executable,
        "platform": platform.platform(),
        "processor": platform.processor(),
        "dependency_file": (
            str(dep_path.relative_to(root)) if dep_path is not None else None
        ),
        "dependency_hash": dep_hash,
        "free_disk_gb": round(free_gb, 2),
        "frozen_files_exist": exists,
        "frozen_file_hashes": hashes,
        "protocol_git_commit": protocol.get("git_commit"),
        "protocol_authorization_status": protocol.get("authorization_status"),
        "protocol_execution_status": protocol.get("execution_status"),
        "protocol_config_hash": hashes.get("protocol"),
        "manifest_git_commit": frozen.get("git_commit"),
        "manifest_e1_config_hash": frozen.get("e1_config_hash"),
        "selected_baseline": baseline.get("selected_baseline"),
        "selected_baseline_hash": hashes.get("baseline"),
        "selected_baseline_git_commit": baseline.get("git_commit"),
        "data_hash": baseline.get("data_hash") or trace_manifest.get("dataset_hash"),
        "target_group_payload_hash": protocol.get("target_group_payload_hash"),
        "target_group_file_hash": protocol.get("target_group_file_hash"),
        "client_mapping_payload_hash": protocol.get("client_mapping_payload_hash"),
        "client_mapping_file_hash": protocol.get("client_mapping_file_hash"),
        "pi_target_hash": protocol.get("pi_target_hash"),
        "pi_manifest": {
            "client_stratum_target_mass_hash": pi_manifest.get(
                "client_stratum_target_mass_hash"
            ),
            "pi_target_file_hash": pi_manifest.get("pi_target_file_hash")
            or pi_manifest.get("pi_target_hash"),
            "atomic_target_weight_hash": pi_manifest.get("atomic_target_weight_hash"),
        },
        "event_traces": trace_rows,
        "event_trace_manifest_git_commit": trace_manifest.get("git_commit"),
        "local_steps_in_frozen_protocol": "local_steps" in protocol,
        "local_steps_runtime_default_in_authorized_code": 2,
        "methods": list(METHODS),
        "seeds": list(SEEDS),
        "formal_runs_started": False,
        "formal_runs_completed": 0,
        "formal_runs_passed": 0,
        "failed_runs": [],
        "resume_count": 0,
        "r4_test_summary_present": test_summary_path.exists(),
        "r4_test_all_exit_zero": bool(test_summary.get("all_exit_zero")),
        "r4_test_failed": test_summary.get("failed"),
        "blockers": blockers,
        "authorization_note": (
            "Teacher selected stop_blocked; no formal 25-run execution started."
        ),
        "e2_e9_status": "NOT_STARTED",
        "estimated_total_runtime_note": (
            "Not estimated for formal execution because preflight blocked before "
            "Phase A; R4 one-seed five-method 20-window smoke completed previously."
        ),
    }


def write_preflight_md(root: Path, preflight: dict) -> Path:
    path = root / "docs/reports/E1_FORMAL_R1_PREFLIGHT.md"
    path.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "# E1-FORMAL-R1 Preflight",
        "",
        f"Generated (UTC): {preflight['generated_at']}",
        f"Status: **{preflight['status']}**",
        "",
        "## Identity",
        f"- Required commit: `{preflight['required_commit']}`",
        f"- HEAD commit: `{preflight['head_commit']}`",
        f"- Git clean: `{preflight['git_clean']}`",
        f"- Python: `{preflight['python_version']}`",
        f"- Platform: `{preflight['platform']}`",
        f"- Processor: `{preflight['processor']}`",
        f"- Dependency file: `{preflight['dependency_file']}`",
        f"- Dependency hash: `{preflight['dependency_hash']}`",
        f"- Free disk (GB): `{preflight['free_disk_gb']}`",
        "",
        "## Frozen Protocol",
        f"- Protocol git_commit: `{preflight['protocol_git_commit']}`",
        f"- Manifest git_commit: `{preflight['manifest_git_commit']}`",
        f"- Protocol config hash: `{preflight['protocol_config_hash']}`",
        f"- Selected baseline: `{preflight['selected_baseline']}`",
        f"- Selected baseline hash: `{preflight['selected_baseline_hash']}`",
        f"- Data hash: `{preflight['data_hash']}`",
        f"- Target-group payload/file: "
        f"`{preflight['target_group_payload_hash']}` / "
        f"`{preflight['target_group_file_hash']}`",
        f"- Client-mapping payload/file: "
        f"`{preflight['client_mapping_payload_hash']}` / "
        f"`{preflight['client_mapping_file_hash']}`",
        f"- Pi-target hash: `{preflight['pi_target_hash']}`",
        f"- Local steps in frozen protocol: "
        f"`{preflight['local_steps_in_frozen_protocol']}`",
        f"- Authorized-code runtime default local_steps: "
        f"`{preflight['local_steps_runtime_default_in_authorized_code']}`",
        "",
        "## EventTraces",
    ]
    for row in preflight["event_traces"]:
        lines.append(
            f"- seed {row['seed']}: events={row['events_exists']}, "
            f"audit_pass={row['audit_pass']}, hash_match={row['hash_match']}, "
            f"generation_commit=`{row['generation_git_commit']}`, "
            f"hash=`{row['actual_hash']}`"
        )
    lines.extend(
        [
            "",
            "## Hard Gates",
            f"- Frozen files exist: `{all(preflight['frozen_files_exist'].values())}`",
            f"- R4 test summary present: `{preflight['r4_test_summary_present']}`",
            f"- R4 test all_exit_zero: `{preflight['r4_test_all_exit_zero']}`",
            f"- Formal runs started: `{preflight['formal_runs_started']}`",
            "",
            "## Blockers",
        ]
    )
    if preflight["blockers"]:
        lines.extend(f"- {item}" for item in preflight["blockers"])
    else:
        lines.append("- none")
    lines.extend(
        [
            "",
            "## Decision",
            "",
            "E1-FORMAL-R1 = BLOCKED. No formal 5x5x100 runs were started.",
            "E2-E9 have NOT started.",
            "",
            preflight["authorization_note"],
            "",
        ]
    )
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def write_report(root: Path, preflight: dict) -> tuple[Path, Path, Path]:
    deliverables = root / "deliverables"
    deliverables.mkdir(parents=True, exist_ok=True)
    sections = [
        (
            "Executive Summary",
            "E1-FORMAL-R1 status is BLOCKED. Formal 25-run execution was not "
            "started after preflight identity conflicts. Teacher selected "
            "stop_blocked.",
        ),
        (
            "Frozen Protocol and Formal Authorization",
            "Authorized code commit is "
            f"`{REQUIRED_COMMIT}`. Frozen protocol/manifest still record "
            f"`{preflight['protocol_git_commit']}` / "
            f"`{preflight['manifest_git_commit']}`. "
            "Authorization status in protocol file remains "
            f"`{preflight['protocol_authorization_status']}`; "
            f"execution_status=`{preflight['protocol_execution_status']}`.",
        ),
        (
            "Git and Experiment Identity",
            f"HEAD=`{preflight['head_commit']}`; git_clean="
            f"`{preflight['git_clean']}`; selected_baseline="
            f"`{preflight['selected_baseline']}`; "
            f"protocol_hash=`{preflight['protocol_config_hash']}`; "
            f"data_hash=`{preflight['data_hash']}`.",
        ),
        (
            "Completion Matrix: 5 Methods × 5 Seeds",
            "0/25 formal runs completed. Matrix is empty by design because "
            "preflight blocked Phase A.",
        ),
        (
            "Per-Run Hard Gates",
            "Not evaluated. No formal run directories were created under "
            "outputs/runs/E1_FORMAL_*.",
        ),
        (
            "Target and Arrival Risk Results",
            "Not available. Formal predictions and arrival weights were not "
            "generated.",
        ),
        (
            "No-Harm Analysis",
            "Not executed. No-harm statistics require 25 successful paired runs.",
        ),
        (
            "Paired Degradation by Seed",
            "Not available.",
        ),
        (
            "One-Sided 95% Upper Bound",
            "Not available.",
        ),
        (
            "Wilcoxon and Holm Results",
            "Not executed.",
        ),
        (
            "Group Risk and Misalignment Gap",
            "Not available.",
        ),
        (
            "ESS and Clipping Safety",
            "Not evaluated for formal 100-window runs. Prior R4 validation "
            "safety remains historical entry evidence only.",
        ),
        (
            "Stage-2 Attempt Semantics",
            "Not re-audited on formal runs. Blocker is freeze identity, not "
            "attempt semantics.",
        ),
        (
            "Arrival Support Audit",
            "Not re-audited on formal runs.",
        ),
        (
            "Solver Reliability",
            "Not evaluated on formal runs.",
        ),
        (
            "Runtime and Communication",
            "No formal runtime collected.",
        ),
        (
            "Failed/Resumed Runs",
            "failed_runs=[]; resume_count=0; formal_runs_started=false.",
        ),
        (
            "E1-G1 to E1-G6",
            "G1 completeness FAIL (0/25). G2 identity FAIL due frozen-commit "
            "mismatch and missing local_steps. G3-G6 not reached.",
        ),
        (
            "Limitations",
            "Frozen protocol git_commit and FROZEN_CONFIG_MANIFEST.git_commit "
            "do not equal the teacher-required formal commit. Frozen protocol "
            "also omits local_steps. Teacher chose to keep the experiment "
            "blocked rather than authorize an interpretation or corrected freeze.",
        ),
        (
            "E1 Final Decision",
            "E1-FORMAL-R1 = BLOCKED. E1 remains not formally executed.",
        ),
        (
            "Next-Step Recommendation",
            "Teacher should either (1) issue corrected frozen protocol/manifest "
            "bound to "
            f"`{REQUIRED_COMMIT}` with explicit local_steps, or "
            "(2) provide written authorization clarifying commit authority and "
            "local_steps=2. Do not start E2-E9.",
        ),
    ]
    md = ["# E1-FORMAL-R1 Report", ""]
    for index, (title, body) in enumerate(sections, 1):
        md.extend([f"## {index}. {title}", body, ""])
    md.extend(
        [
            "## Blockers",
            *[f"- {item}" for item in preflight["blockers"]],
            "",
            "Formal five-seed E1 has NOT been executed.",
            "E2-E9 have NOT started.",
            "",
        ]
    )
    md_path = deliverables / "E1_FORMAL_REPORT.md"
    md_path.write_text("\n".join(md), encoding="utf-8")

    doc = Document()
    doc.add_heading("E1-FORMAL-R1 Report", 0)
    doc.add_paragraph(f"Status: {preflight['status']}")
    doc.add_paragraph(f"Generated (UTC): {preflight['generated_at']}")
    for index, (title, body) in enumerate(sections, 1):
        doc.add_heading(f"{index}. {title}", level=1)
        doc.add_paragraph(body)
    doc.add_heading("Blockers", level=1)
    for item in preflight["blockers"]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_paragraph(
        "Formal five-seed E1 has NOT been executed. E2-E9 have NOT started."
    )
    docx_path = deliverables / "E1_FORMAL_REPORT.docx"
    doc.save(docx_path)

    readme = deliverables / "E1_FORMAL_SUBMISSION_README.txt"
    readme.write_text(
        "\n".join(
            [
                "E1-FORMAL-R1 blocked evidence package",
                f"Required commit: {REQUIRED_COMMIT}",
                f"HEAD commit: {preflight['head_commit']}",
                f"Status: {preflight['status']}",
                "Formal 25-run execution was not started.",
                "E2-E9 have NOT started.",
                "See E1_FORMAL_REPORT.docx and docs/reports/E1_FORMAL_R1_PREFLIGHT.md.",
                "",
            ]
        ),
        encoding="utf-8",
    )
    return md_path, docx_path, readme


def write_git_evidence(root: Path) -> Path:
    out = root / "outputs/evidence_formal_r1/git"
    out.mkdir(parents=True, exist_ok=True)
    (out / "GIT_HEAD.txt").write_text(
        _git(root, "rev-parse", "HEAD"), encoding="utf-8"
    )
    (out / "GIT_STATUS.txt").write_text(
        _git(root, "status", "--porcelain"), encoding="utf-8"
    )
    (out / "GIT_LOG.txt").write_text(
        _git(root, "log", "-20", "--oneline"), encoding="utf-8"
    )
    (out / "GIT_DIFF_SUMMARY.txt").write_text(
        _git(root, "diff", "--stat"), encoding="utf-8"
    )
    return out


def package_zip(
    root: Path,
    preflight_path: Path,
    preflight_md: Path,
    md_path: Path,
    docx_path: Path,
    readme: Path,
    git_dir: Path,
) -> Path:
    zip_path = root / "deliverables/RAVEN_MCS_E1_FORMAL_EVIDENCE.zip"
    include_dirs = [
        root / "configs/frozen",
        root / "outputs/event_traces",
        root / "outputs/audits",
        root / "outputs/evidence_formal_r1",
        root / "logs",
        root / "docs/reports",
    ]
    include_files = [
        root / "STATUS.md",
        root / "ISSUES.md",
        root / "CHANGELOG.md",
        root / "docs/FORMULA_TO_CODE_MAP.md",
        root / "docs/DATA_DICTIONARY.md",
        preflight_path,
        preflight_md,
        md_path,
        docx_path,
        readme,
    ]
    tracked = _git(root, "ls-files").splitlines()
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
        seen: set[str] = set()

        def add(path: Path) -> None:
            if not path.exists() or not path.is_file():
                return
            rel = path.relative_to(root).as_posix()
            if rel in seen:
                return
            archive.write(path, rel)
            seen.add(rel)

        for relative in tracked:
            add(root / relative)
        for directory in include_dirs:
            if not directory.exists():
                continue
            for path in directory.rglob("*"):
                if path.is_file():
                    add(path)
        for path in include_files:
            add(path)
        for path in git_dir.rglob("*"):
            if path.is_file():
                add(path)
    return zip_path


def main() -> int:
    root = Path(__file__).resolve().parents[1]
    preflight = collect_preflight(root)
    audits = root / "outputs/audits"
    audits.mkdir(parents=True, exist_ok=True)
    preflight_path = audits / "e1_formal_r1_preflight.json"
    preflight_path.write_text(json.dumps(preflight, indent=2), encoding="utf-8")
    preflight_md = write_preflight_md(root, preflight)
    md_path, docx_path, readme = write_report(root, preflight)
    git_dir = write_git_evidence(root)
    zip_path = package_zip(
        root,
        preflight_path,
        preflight_md,
        md_path,
        docx_path,
        readme,
        git_dir,
    )
    hashes = {
        "preflight_json": sha256_file(preflight_path),
        "preflight_md": sha256_file(preflight_md),
        "report_md": sha256_file(md_path),
        "report_docx": sha256_file(docx_path),
        "readme": sha256_file(readme),
        "evidence_zip": sha256_file(zip_path),
        "status": preflight["status"],
        "head_commit": preflight["head_commit"],
    }
    (root / "deliverables/E1_FORMAL_DELIVERABLE_HASHES.json").write_text(
        json.dumps(hashes, indent=2), encoding="utf-8"
    )
    print(json.dumps({
        "status": preflight["status"],
        "blockers": preflight["blockers"],
        "report": str(docx_path),
        "zip": str(zip_path),
        "zip_sha256": hashes["evidence_zip"],
        "formal_runs_started": False,
    }, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
