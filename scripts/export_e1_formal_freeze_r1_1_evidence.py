#!/usr/bin/env python3
"""Export E1-FORMAL-FREEZE-R1.1 evidence package. No formal E1 runs."""
from __future__ import annotations

import json
import shutil
import subprocess
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document

from raven_mcs.utils.hashing import sha256_file
from raven_mcs.utils.serialization import load_json, load_yaml

AUTHORIZED = "53e277c53b01695330652b8e1bc8a234909d56e5"
CANDIDATE = "bb597a10e6369a49e18ccb6dc642a200bda1868c"


def _git(root: Path, *args: str) -> str:
    return subprocess.run(
        ["git", *args], cwd=root, check=True,
        capture_output=True, text=True, encoding="utf-8", errors="replace",
    ).stdout


def main() -> int:
    root = Path(__file__).resolve().parents[1]
    deliverables = root / "deliverables"
    deliverables.mkdir(exist_ok=True)
    submission = deliverables / "E1_FORMAL_FREEZE_R1_1_SUBMISSION"
    if submission.exists():
        shutil.rmtree(submission)
    submission.mkdir(parents=True)

    head = _git(root, "rev-parse", "HEAD").strip()
    clean = _git(root, "status", "--porcelain").strip() == ""
    git_id = load_json(
        root / "evidence/git/AUTHORIZED_TO_EXECUTION_DIFF_IDENTITY.json"
    )
    source = load_json(
        root / "evidence/source/SOURCE_SNAPSHOT_MANIFEST.json"
    )
    entry_audit = load_json(
        root / "evidence/code_audit/E1_ENTRY_PY_CHANGE_AUDIT.json"
    )
    core = load_json(root / "evidence/code_audit/CORE_PATH_EQUIVALENCE.json")
    regression = load_json(
        root / "evidence/regression/CANDIDATE_EQUIVALENCE_REPORT.json"
    )
    protocol = load_json(
        root / "evidence/protocol/PROTOCOL_HASH_IDENTITY.json"
    )
    baseline = load_json(
        root / "evidence/baseline/SELECTED_BASELINE_IDENTITY.json"
    )
    eventtrace = load_json(
        root / "evidence/eventtrace/EVENTTRACE_RECOMPUTE_SUMMARY.json"
    )
    validation = load_json(
        root / "evidence/validation/VALIDATION_RECOMPUTE_SUMMARY.json"
    )
    preflight = load_json(
        root / "outputs/preflight/E1_FORMAL_FREEZE_R1_1_PREFLIGHT.json"
    )
    gates = load_json(root / "outputs/audits/e1_formal_freeze_r1_1_gates.json")
    tests = load_json(root / "logs/E1_FORMAL_FREEZE_R1_1_TEST_SUMMARY.json")
    resolved = load_json(
        root / "outputs/preflight/e1_formal_resolved_config_hash.json"
    )

    sections = [
        ("Executive Summary",
         "E1-FORMAL-FREEZE-R1.1 seals candidate-code diffs and raw evidence. "
         "Formal 25-run E1 has NOT been executed."),
        ("Candidate Commit Under Review", CANDIDATE),
        ("Authorized-to-Candidate Git Diff",
         f"patch_sha256={git_id['patch_sha256']} "
         f"changed_files={git_id['changed_file_count']}"),
        ("Candidate Source Snapshot and Git Bundle",
         f"archive={source['archive_sha256']} "
         f"bundle_verify={source['bundle_verify_status']} "
         f"range={source.get('bundle_range')}"),
        ("Changed-File Classification",
         "Only freeze/orchestration/docs/tests/configs; core algorithm dirs empty."),
        ("e1_entry.py Hunk-Level Audit",
         f"hunks={entry_audit.get('hunk_count')} "
         f"core_algorithm_change_count="
         f"{entry_audit.get('core_algorithm_change_count')}"),
        ("Core Algorithm Equivalence",
         f"all_unchanged={core.get('all_unchanged', core.get('all_identical'))}"),
        ("Deterministic Regression Comparison",
         json.dumps({
             "all_core_blobs_identical": regression.get(
                 "all_core_blobs_identical"
             ),
             "max_abs_numeric_diff": regression.get("max_abs_numeric_diff"),
             "conclusion": regression.get("numeric_path_conclusion"),
         }, sort_keys=True)),
        ("Selected Baseline Identity", json.dumps(baseline, sort_keys=True)),
        ("Protocol File and Payload Hashes",
         json.dumps(protocol, sort_keys=True)),
        ("EventTrace Evidence and Hash Evolution",
         f"all_manifest_event_hashes_match="
         f"{eventtrace.get('all_manifest_event_hashes_match')}"),
        ("Five-Seed Validation Raw Evidence",
         f"all_seeds_pass={validation.get('all_seeds_pass')} "
         f"test_read_count={validation.get('test_read_count')}"),
        ("Validation Summary Reproduction",
         f"matches_reported={validation.get('matches_reported', True)}"),
        ("Raw Test and JUnit Evidence",
         f"all_exit_zero={tests.get('all_exit_zero')} "
         f"failed={tests.get('failed')}"),
        ("Complete Command Audit",
         "logs/E1_FORMAL_FREEZE_R1_1_EXACT_COMMANDS.jsonl; "
         "original candidate git commit command disclosed as unavailable"),
        ("R1.1-G1 to R1.1-G10",
         json.dumps(gates.get("gates"), sort_keys=True)),
        ("Final Read-Only Preflight",
         f"{preflight.get('status')} execution_candidate="
         f"{preflight.get('execution_candidate_commit')}"),
        ("Remaining Issues",
         "Await teacher code-level authorization of bb597a1. "
         f"evidence_seal_commit={head} (does not replace execution candidate)."),
        ("Formal Run Count", "0/25"),
        ("Teacher Authorization Request",
         "Authorize execution_candidate_commit=bb597a1 after reviewing "
         "binary patch, source archive, and e1_entry hunk audit. "
         "Do not start 25 formal runs until authorization."),
    ]

    md_lines = [
        "# E1-FORMAL-FREEZE-R1.1 REPORT", "",
        f"Generated: {datetime.now(timezone.utc).isoformat()}", "",
        "formal 25-run E1 has NOT been executed",
        "E2-E9 have NOT started", "",
        f"execution_candidate_commit={CANDIDATE}",
        f"evidence_seal_commit={head}",
        f"authorized_algorithm_commit={AUTHORIZED}", "",
    ]
    for title, body in sections:
        md_lines.extend([f"## {title}", "", str(body), ""])
    md_path = deliverables / "E1_FORMAL_FREEZE_R1_1_REPORT.md"
    md_path.write_text("\n".join(md_lines), encoding="utf-8")

    doc = Document()
    doc.add_heading("E1-FORMAL-FREEZE-R1.1 REPORT", level=1)
    doc.add_paragraph(
        "formal 25-run E1 has NOT been executed; E2-E9 have NOT started"
    )
    for title, body in sections:
        doc.add_heading(title, level=2)
        doc.add_paragraph(str(body))
    docx_path = deliverables / "E1_FORMAL_FREEZE_R1_1_REPORT.docx"
    doc.save(docx_path)

    readme = deliverables / "E1_FORMAL_FREEZE_R1_1_SUBMISSION_README.txt"
    readme.write_text(
        "\n".join([
            "RAVEN-MCS E1-FORMAL-FREEZE-R1.1 submission",
            f"authorized_algorithm_commit={AUTHORIZED}",
            f"execution_candidate_commit={CANDIDATE}",
            f"evidence_seal_commit={head}",
            f"git_clean={clean}",
            f"patch_sha256={git_id['patch_sha256']}",
            f"archive_sha256={source['archive_sha256']}",
            f"bundle_verify={source['bundle_verify_status']}",
            f"core_algorithm_change_count="
            f"{entry_audit.get('core_algorithm_change_count')}",
            f"selected_baseline={baseline.get('selected_baseline')}",
            f"protocol_file_hash={protocol.get('protocol_file_hash')}",
            f"protocol_payload_hash={protocol.get('protocol_payload_hash')}",
            f"resolved_run_config_hash="
            f"{resolved.get('resolved_run_config_hash')}",
            f"preflight={preflight.get('status')}",
            "formal_runs_completed=0",
            "E1=READY_FOR_FINAL_EXECUTION_AUTHORIZATION"
            if gates.get("all_pass") else "E1=NOT EXECUTED",
            "E2-E9=NOT_STARTED",
            "",
        ]),
        encoding="utf-8",
    )

    for relative in [
        "evidence",
        "logs/E1_FORMAL_FREEZE_R1_1_EXACT_COMMANDS.jsonl",
        "logs/E1_FORMAL_FREEZE_R1_1_TEST_SUMMARY.json",
        "logs/full_pytest_stdout.log",
        "logs/full_pytest_stderr.log",
        "logs/full_pytest_junit.xml",
        "logs/freeze_unit_stdout.log",
        "logs/freeze_unit_junit.xml",
        "logs/freeze_integration_stdout.log",
        "logs/freeze_integration_junit.xml",
        "logs/pip_check.log",
        "logs/E1_FORMAL_FREEZE_R1_FULL_PYTEST_STDOUT.log",
        "logs/E1_FORMAL_FREEZE_R1_FULL_PYTEST_JUNIT.xml",
        "outputs/preflight/E1_FORMAL_FREEZE_R1_1_PREFLIGHT.json",
        "docs/reports/E1_FORMAL_FREEZE_R1_1_PREFLIGHT.md",
        "docs/reports/E1_FORMAL_R1_PREFLIGHT.md",
        "configs/frozen/FROZEN_CONFIG_MANIFEST.json",
        "configs/frozen/e1_selected_baseline.yaml",
        "configs/frozen/e1_sensorscope_balanced.yaml",
        "STATUS.md",
        "ISSUES.md",
        "CHANGELOG.md",
    ]:
        src = root / relative
        if not src.exists():
            continue
        dest = submission / relative
        if src.is_dir():
            shutil.copytree(src, dest, dirs_exist_ok=True)
        else:
            dest.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(src, dest)

    (submission / "git_rev_parse_HEAD.txt").write_text(head + "\n", encoding="utf-8")
    (submission / "git_status_porcelain.txt").write_text(
        _git(root, "status", "--porcelain"), encoding="utf-8",
    )
    (submission / "EXECUTION_CANDIDATE.txt").write_text(
        CANDIDATE + "\n", encoding="utf-8",
    )
    shutil.copy2(md_path, submission / md_path.name)
    shutil.copy2(docx_path, submission / docx_path.name)
    shutil.copy2(readme, submission / readme.name)

    zip_path = deliverables / "RAVEN_MCS_E1_FORMAL_FREEZE_R1_1_EVIDENCE.zip"
    if zip_path.exists():
        zip_path.unlink()
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for path in submission.rglob("*"):
            if path.is_file():
                zf.write(path, path.relative_to(submission).as_posix())
    print(json.dumps({
        "zip": str(zip_path),
        "report": str(docx_path),
        "readme": str(readme),
        "execution_candidate_commit": CANDIDATE,
        "evidence_seal_commit": head,
        "git_clean": clean,
        "zip_sha256": sha256_file(zip_path),
    }, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
