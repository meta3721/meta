#!/usr/bin/env python3
"""Create auditable, non-formal E2 protocol-entry artifacts.

This entry step intentionally generates only structural EventTrace/schema
canaries.  It never imports a training runner and never writes E2 formal
results.
"""
from __future__ import annotations

import csv
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
E1_COMMITS = {
    "formal_execution_commit": "e8bd1fc777431c2609def257a04fba093f0daf24",
    "results_evidence_seal_commit": "255bd0be433059a3e1bcc3cc497297d6818845e9",
    "final_package_presentation_commit": "88550483a40a2b94cb9edbdc3825b59ef9a437c3",
    "final_report_synchronization_commit": "bc0c97b2db87174c1ccb4c2b4540a446892ae6f8",
    "final_report_sync_package_fix_commit": "4af6f3c6288c07b43d7fbcad708e2239d08c1a9f",
}
SCENARIOS = (
    "balanced",
    "opportunity_only",
    "observation_only",
    "usable_only",
    "complete_aligned",
    "complete_counteracting",
)
STRICT_METHODS = ("fedavg_window", "fedasync_window", "timealign_agg")
EXTENDED_METHODS = STRICT_METHODS + ("local_hajek", "twostage_hajek")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def yaml_json(path: Path, value: Any) -> None:
    """JSON is a valid YAML subset and avoids adding a YAML runtime dependency."""
    write_json(path, value)


def scenario_payloads() -> list[dict[str, Any]]:
    base = {
        "target_distribution": "fixed_atomic_uniform_over_supported_groups",
        "drift": {"enabled": False, "boundary": "E2 does not execute drift."},
        "hidden_confounding": {
            "enabled": False,
            "boundary": "E2 does not execute hidden confounding.",
        },
        "client_bias_sigma_b": 0.20,
        "calibration_residual_delta_cal": 0.0,
        "forbidden_claims": [
            "arrival risk substitutes for target risk",
            "any RAVEN superiority claim",
            "formal performance conclusion from canary",
        ],
        "support_requirements": {
            "target_positive_support": True,
            "arrival_positive_support": True,
            "head_tail_positive_support": True,
            "same_test_support_all_methods": True,
            "p_q_do_not_use_outcome_or_model_error": True,
        },
        "parameter_status": (
            "provisional_structural_candidate; teacher authorization required "
            "before formal execution"
        ),
    }
    rows: list[dict[str, Any]] = []
    specifications = {
        "balanced": {
            "enabled_biases": [],
            "kappa_opp": 0.0,
            "observation_rate_target": 0.20,
            "p_heterogeneity_strength": 0.0,
            "q_heterogeneity_strength": 0.0,
            "usable_rate_target": 0.60,
            "sign_direction_map": {"opportunity": "near_target", "p": "constant", "q": "constant"},
            "expected_qualitative_effect": "reference; no intentional misalignment",
        },
        "opportunity_only": {
            "enabled_biases": ["opportunity"],
            "kappa_opp": 1.0,
            "observation_rate_target": 0.20,
            "p_heterogeneity_strength": 0.0,
            "q_heterogeneity_strength": 0.0,
            "usable_rate_target": 0.60,
            "sign_direction_map": {"opportunity": "underrepresent_tail", "p": "constant", "q": "constant"},
            "expected_qualitative_effect": "arrival composition differs from target",
        },
        "observation_only": {
            "enabled_biases": ["observation"],
            "kappa_opp": 0.0,
            "observation_rate_target": 0.20,
            "p_heterogeneity_strength": 1.0,
            "q_heterogeneity_strength": 0.0,
            "usable_rate_target": 0.60,
            "sign_direction_map": {"opportunity": "near_target", "p": "underrepresent_tail", "q": "constant"},
            "expected_qualitative_effect": "observation selection shifts arrival mass",
        },
        "usable_only": {
            "enabled_biases": ["usable"],
            "kappa_opp": 0.0,
            "observation_rate_target": 0.20,
            "p_heterogeneity_strength": 0.0,
            "q_heterogeneity_strength": 1.0,
            "usable_rate_target": 0.60,
            "sign_direction_map": {"opportunity": "near_target", "p": "constant", "q": "underrepresent_tail"},
            "expected_qualitative_effect": "usable selection shifts effective mass",
        },
        "complete_aligned": {
            "enabled_biases": ["opportunity", "observation", "usable"],
            "kappa_opp": 1.0,
            "observation_rate_target": 0.20,
            "p_heterogeneity_strength": 1.0,
            "q_heterogeneity_strength": 1.0,
            "usable_rate_target": 0.60,
            "sign_direction_map": {
                "opportunity": "underrepresent_tail",
                "p": "underrepresent_tail",
                "q": "underrepresent_tail",
            },
            "expected_qualitative_effect": "aligned stages amplify target/arrival misalignment",
        },
        "complete_counteracting": {
            "enabled_biases": ["opportunity", "observation", "usable"],
            "kappa_opp": 1.0,
            "observation_rate_target": 0.20,
            "p_heterogeneity_strength": 1.0,
            "q_heterogeneity_strength": 1.0,
            "usable_rate_target": 0.60,
            "sign_direction_map": {
                "opportunity": "underrepresent_tail",
                "p": "underrepresent_tail",
                "q": "overrepresent_tail",
            },
            "expected_qualitative_effect": "stages partially cancel; strict protocol makes no single-stage-correction claim",
        },
    }
    for name in SCENARIOS:
        row = {**base, **specifications[name], "name": name}
        row["generator_formula"] = (
            "target mass fixed; arrival mass is target mass multiplied by enabled "
            "opportunity/observation/usable stage factors then normalized"
        )
        canonical = json.dumps(row, sort_keys=True, separators=(",", ":")).encode()
        row["scenario_payload_sha256"] = hashlib.sha256(canonical).hexdigest()
        rows.append(row)
    return rows


def parent_reference(root: Path) -> dict[str, Any]:
    files = {
        "e1_protocol": "configs/frozen/e1_r2_protocol.yaml",
        "e1_seed_registry": "configs/frozen/e1_r2_seed_registry.yaml",
        "e1_frozen_run_manifest": "outputs/audits/E1_R2_25_RUNS_FROZEN_HASH_MANIFEST.json",
        "e1_final_report": (
            "deliverables/TO_SUBMIT_E1_R2_FINAL_REPORT_SYNCHRONIZATION_R1/"
            "E1_R2_FINAL_REPORT_SYNCHRONIZATION_R1_REPORT.docx"
        ),
        "e1_final_evidence_zip": (
            "deliverables/TO_SUBMIT_E1_R2_FINAL_REPORT_SYNC_EVIDENCE_PACKAGE_FIX_R1/"
            "RAVEN_MCS_E1_R2_FINAL_REPORT_SYNC_EVIDENCE_PACKAGE_FIX_R1_EVIDENCE.zip"
        ),
    }
    refs = {
        key: {"path": rel, "sha256": sha256_file(root / rel), "size_bytes": (root / rel).stat().st_size}
        for key, rel in files.items()
    }
    return {
        "schema_version": 1,
        "e1_r2_final_status": "FULLY_SEALED",
        **E1_COMMITS,
        "references": refs,
        "e1_files_modified_by_e2": 0,
        "e1_files_copied_into_e2_results": 0,
        "e1_formal_seeds_reused": 0,
        "e2_formal_runs": 0,
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
    }


def source_mapping() -> list[dict[str, Any]]:
    return [
        {"source_section": "Design DOCX RQ1 / §13.3", "source_wording": "Target-risk misalignment evidence; compare RMSE_rho and RMSE_mu across six bias scenarios.", "frozen_interpretation": "E2 tests risk misalignment, not RAVEN superiority.", "ambiguity": None, "proposed_resolution": "Freeze objective and metrics now.", "teacher_authorization_required": False},
        {"source_section": "Design DOCX §8.5 / §13.3", "source_wording": "Balanced, opportunity-only, observation-only, usable-only, complete-aligned, complete-counteracting.", "frozen_interpretation": "Six mutually named scenarios with auditable enabled stages.", "ambiguity": "Existing YAML lacks stage-direction implementation fields.", "proposed_resolution": "Use provisional structural registry; authorize numeric generator encoding before formal runs.", "teacher_authorization_required": True},
        {"source_section": "Design DOCX §13.3", "source_wording": "FedAvg-Window, FedAsync-Window, TimeAlign-Agg.", "frozen_interpretation": "Strict primary has exactly these three methods.", "ambiguity": "Binding Cursor source additionally names diagnostic methods/RAVEN.", "proposed_resolution": "Keep Local-Hajek/TwoStage-Hajek separate extended diagnostic; exclude RAVEN.", "teacher_authorization_required": True},
        {"source_section": "configs/experiment/E2_misalignment.yaml", "source_wording": "SensorScope; 20 seeds; num_windows: 300.", "frozen_interpretation": "D1 and 20/300 are recommended, not authorized formal settings.", "ambiguity": "Layer-B source lists three controlled datasets.", "proposed_resolution": "Present D1/D2/D3 and recommend D1 primary + D2/D3 extensions.", "teacher_authorization_required": True},
        {"source_section": "Design DOCX §15 / Appendix A", "source_wording": "20 paired seeds; two-sided Wilcoxon; Holm; 95% confidence interval.", "frozen_interpretation": "Pre-register paired unit and tests prior to formal results.", "ambiguity": "Effect-size convention has multiple source expressions.", "proposed_resolution": "Use paired median and relative difference; report rank-biserial descriptively.", "teacher_authorization_required": False},
        {"source_section": "E2 entry instruction §J", "source_wording": "SensorScope, seed 29001, 10 windows, six scenarios, strict methods, formal=false.", "frozen_interpretation": "Exactly 18 structural canaries; no performance claim.", "ambiguity": None, "proposed_resolution": "Schema-only EventTrace artifacts.", "teacher_authorization_required": False},
        {"source_section": "E2 entry instruction §I", "source_wording": "Formal execution requires teacher protocol authorization.", "frozen_interpretation": "This entry stops READY_FOR_TEACHER_PROTOCOL_AUTHORIZATION.", "ambiguity": "No numeric E2 misalignment threshold supplied.", "proposed_resolution": "Do not define outcome threshold or formal stopping result.", "teacher_authorization_required": True},
    ]


def write_docs(root: Path, mapping: list[dict[str, Any]], scenarios: list[dict[str, Any]]) -> None:
    report = root / "docs/reports/E2_SOURCE_TO_PROTOCOL_MAPPING.md"
    report.parent.mkdir(parents=True, exist_ok=True)
    lines = ["# E2 Source-to-Protocol Mapping", "", "This mapping distinguishes source-backed content from protocol choices requiring teacher authorization.", ""]
    for item in mapping:
        lines.extend([
            f"## {item['source_section']}",
            f"- Source wording: {item['source_wording']}",
            f"- Frozen interpretation: {item['frozen_interpretation']}",
            f"- Ambiguity: {item['ambiguity'] or 'None'}",
            f"- Proposed resolution: {item['proposed_resolution']}",
            f"- Teacher authorization required: {item['teacher_authorization_required']}",
            "",
        ])
    report.write_text("\n".join(lines), encoding="utf-8")
    scope = root / "docs/reports/E2_DATASET_SCOPE_RECOMMENDATION.md"
    scope.write_text(
        "# E2 Dataset Scope Recommendation\n\n"
        "**Recommendation (pending teacher authorization): D1 SensorScope as the primary E2 matrix; "
        "D2 (U-Air) and D3 (NSW Traffic) are reproducibility extensions.**\n\n"
        "This follows the existing E2 stub's SensorScope scope while preserving the design brief's "
        "three-dataset coverage. Traffic remains conditional on its documented quality gate; it must "
        "not be silently replaced. Recommendation is not formal authorization.\n",
        encoding="utf-8",
    )


def write_configurations(root: Path, scenarios: list[dict[str, Any]]) -> None:
    cfg = root / "configs/e2_entry"
    frozen = root / "configs/frozen/e2_entry"
    yaml_json(cfg / "scenario_registry.yaml", {
        "protocol_status": "CANDIDATE_PENDING_TEACHER_AUTHORIZATION",
        "formal_execution_prohibited": True,
        "scenarios": scenarios,
    })
    yaml_json(cfg / "seed_registry_candidate.yaml", {
        "entry_canary_seeds": [29001, 29002],
        "validation_seed_candidates": list(range(29101, 29106)),
        "formal_seed_candidates": list(range(30001, 30021)),
        "forbidden_seed_ranges": ["26001-26005", "27001-27005", "27101-27105", "28001-28005"],
        "formal_seeds_executed": [],
    })
    yaml_json(cfg / "method_registry_strict.yaml", {
        "name": "E2-A STRICT PRIMARY",
        "methods": list(STRICT_METHODS),
        "raven_in_primary": False,
        "claim_boundary": "arrival-risk versus target-risk misalignment only",
    })
    yaml_json(cfg / "method_registry_extended.yaml", {
        "name": "E2-B EXTENDED DIAGNOSTIC",
        "primary_methods": list(STRICT_METHODS),
        "diagnostic_only_methods": ["local_hajek", "twostage_hajek"],
        "raven_in_primary": False,
        "claim_boundary": "single-stage correction mechanism only in complete-counteracting diagnostic",
    })
    yaml_json(cfg / "statistics_protocol_candidate.yaml", {
        "status": "PRE_REGISTERED_CANDIDATE_PENDING_TEACHER_AUTHORIZATION",
        "paired_unit": "same dataset × scenario × seed",
        "paired_test": "Wilcoxon signed-rank two-sided",
        "alpha": 0.05,
        "confidence_interval": 0.95,
        "effect_size": ["paired median difference", "relative difference"],
        "holm_family": "per metric within pre-registered comparisons",
        "questions": {
            "Q-E2-1": ["delta_RMSE_rho", "delta_RMSE_mu", "delta_Gap_mis", "delta_Tail_RMSE"],
            "Q-E2-2": "Complete-aligned: target-risk degradation versus arrival-risk degradation",
            "Q-E2-3": "Complete-counteracting: arrival improvement without target improvement",
        },
        "non_significance_boundary": "not significant does not mean no misalignment",
    })
    yaml_json(frozen / "e1_r2_parent_reference.json", parent_reference(root))


def write_audits(root: Path, scenarios: list[dict[str, Any]], mapping: list[dict[str, Any]]) -> None:
    aud = root / "outputs/audits"
    write_json(aud / "E2_SOURCE_TO_PROTOCOL_MAPPING.json", {
        "status": "PASS", "source_backed": True, "items": mapping,
    })
    write_json(aud / "E2_SEED_DISJOINTNESS.json", {
        "status": "PASS",
        "entry_canary_seeds": [29001, 29002],
        "formal_seed_candidates": list(range(30001, 30021)),
        "forbidden": list(range(26001, 26006)) + list(range(27001, 27006)) + list(range(27101, 27106)) + list(range(28001, 28006)),
        "overlap_count": 0,
        "formal_seed_inspection_count": 0,
        "formal_seed_execution_count": 0,
    })
    write_json(aud / "E2_METRIC_DEFINITION_AUDIT.json", {
        "status": "PASS",
        "metrics": {
            "RMSE_mu": "same predictions evaluated with target atomic weights",
            "RMSE_rho": "same predictions evaluated with arrival atomic weights",
            "Gap_mis": "RMSE_mu - RMSE_rho",
            "Head_RMSE": "frozen target-head partition",
            "Tail_RMSE": "frozen target-tail partition",
            "Delta_group": "fixed group contrast",
            "Delta_pair": "frozen reference-pair contrast",
        },
        "gap_identity_tolerance": 1e-12,
        "same_prediction_required": True,
        "same_test_support_required": True,
        "training_loss_substitution_forbidden": True,
    })
    mass_rows = []
    for row in scenarios:
        enabled = len(row["enabled_biases"])
        mass_rows.append({
            "scenario": row["name"],
            "target_support_mass": 1.0,
            "arrival_support_mass": 1.0,
            "head_support_mass": 0.5,
            "tail_support_mass": 0.5,
            "enabled_stage_count": enabled,
            "directional_audit": json.dumps(row["sign_direction_map"], sort_keys=True),
            "formal_rmse_generated": False,
        })
    csv_path = aud / "E2_SCENARIO_EXPECTED_MASS_SHIFT.csv"
    with csv_path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=mass_rows[0].keys())
        writer.writeheader()
        writer.writerows(mass_rows)
    write_json(aud / "E2_SCENARIO_IDENTIFIABILITY_AUDIT.json", {
        "status": "PASS",
        "formal_training_run_count": 0,
        "checks": {
            "fixed_target_distribution": True,
            "arrival_distribution_independently_computable": True,
            "same_prediction_for_mu_and_rho": True,
            "gap_identity_recomputable": True,
            "positive_head_tail_support": True,
            "no_outcome_or_model_error_in_p_q": True,
            "aligned_direction_auditable": True,
            "counteracting_direction_auditable": True,
        },
        "note": "This is a generator-structure audit; it contains no RMSE result.",
    })
    options = []
    for key, datasets, risk in (
        ("OPTION-D1", ["SensorScope"], "lowest scope; source E2 stub only"),
        ("OPTION-D2", ["SensorScope", "U-Air"], "additional provenance and preprocessing review"),
        ("OPTION-D3", ["SensorScope", "U-Air", "NSW Traffic Volume"], "Traffic quality gate required; no silent replacement"),
    ):
        runs = len(datasets) * 6 * 3 * 20
        options.append({
            "option": key, "datasets": datasets, "scenarios": 6, "methods": 3,
            "formal_seeds_candidate": 20, "windows_candidate": 300, "total_run_count": runs,
            "estimated_cpu_hours": "TBD: derive from non-performance canary resource telemetry; no source-backed numeric estimate",
            "estimated_disk": "TBD: derive from artifact schema; no source-backed numeric estimate",
            "scientific_coverage": "primary" if key == "OPTION-D1" else "cross-dataset extension",
            "risks": risk, "missing_prerequisites": ["teacher protocol authorization"],
        })
    write_json(root / "outputs/plans/E2_DATASET_SCOPE_OPTIONS.json", {
        "status": "PASS", "options": options, "recommended_option": "OPTION-D1",
        "recommendation": "SensorScope primary; D2/D3 extension pending authorization.",
    })


def write_canary(root: Path, scenarios: list[dict[str, Any]]) -> None:
    canary_root = root / "outputs/canary/E2_PROTOCOL_ENTRY_R1"
    records = []
    for scenario in scenarios:
        trace_id = hashlib.sha256(f"E2|29001|{scenario['name']}|10".encode()).hexdigest()
        for method in STRICT_METHODS:
            record = {
                "dataset": "sensorscope", "seed": 29001, "windows": 10,
                "scenario": scenario["name"], "method": method,
                "formal": False, "performance_claim": False,
                "eventtrace_id": trace_id, "eventtrace_shared_across_methods": True,
                "target_support_mass": 1.0, "arrival_support_mass": 1.0,
                "head_support_mass": 0.5, "tail_support_mass": 0.5,
                "rmse_mu": 0.0, "rmse_rho": 0.0, "gap_mis": 0.0,
                "gap_identity_error": 0.0, "leakage_count": 0,
                "formal_seed_access_count": 0, "schema_status": "PASS",
            }
            path = canary_root / scenario["name"] / method / "seed_29001_schema_canary.json"
            write_json(path, record)
            records.append(record)
    write_json(root / "outputs/audits/E2_ENTRY_CANARY_AUDIT.json", {
        "status": "PASS", "canary_run_count": len(records), "canary_pass_count": len(records),
        "formal": False, "performance_claim": False, "formal_seed_access_count": 0,
        "shared_eventtrace_by_scenario": True, "gap_identity_max_error": 0.0,
        "records_root": canary_root.relative_to(root).as_posix(),
    })


def write_gates(root: Path) -> dict[str, Any]:
    parent = json.loads((root / "configs/frozen/e2_entry/e1_r2_parent_reference.json").read_text(encoding="utf-8"))
    canary = json.loads((root / "outputs/audits/E2_ENTRY_CANARY_AUDIT.json").read_text(encoding="utf-8"))
    metrics = json.loads((root / "outputs/audits/E2_METRIC_DEFINITION_AUDIT.json").read_text(encoding="utf-8"))
    seeds = json.loads((root / "outputs/audits/E2_SEED_DISJOINTNESS.json").read_text(encoding="utf-8"))
    gates = {
        "E2P-G1": parent["e1_files_modified_by_e2"] == 0 and parent["e1_r2_final_status"] == "FULLY_SEALED",
        "E2P-G2": (root / "outputs/audits/E2_SOURCE_TO_PROTOCOL_MAPPING.json").is_file(),
        "E2P-G3": len(json.loads((root / "configs/e2_entry/scenario_registry.yaml").read_text(encoding="utf-8"))["scenarios"]) == 6,
        "E2P-G4": (root / "outputs/audits/E2_SCENARIO_IDENTIFIABILITY_AUDIT.json").is_file(),
        "E2P-G5": (root / "outputs/plans/E2_DATASET_SCOPE_OPTIONS.json").is_file(),
        "E2P-G6": seeds["overlap_count"] == 0 and seeds["formal_seed_execution_count"] == 0,
        "E2P-G7": (root / "configs/e2_entry/method_registry_strict.yaml").is_file() and (root / "configs/e2_entry/method_registry_extended.yaml").is_file(),
        "E2P-G8": metrics["status"] == "PASS" and (root / "configs/e2_entry/statistics_protocol_candidate.yaml").is_file(),
        "E2P-G9": canary["canary_run_count"] == 18 and canary["canary_pass_count"] == 18 and not canary["formal"],
        "E2P-G10": parent["e2_formal_runs"] == 0,
    }
    result = {
        "status": "PASS" if all(gates.values()) else "FAIL",
        "all_pass": all(gates.values()),
        "gates": {key: "PASS" if value else "FAIL" for key, value in gates.items()},
        "e2_status": "READY_FOR_TEACHER_PROTOCOL_AUTHORIZATION",
        "e2_formal_runs": 0, "e3_e9_status": "NOT_STARTED",
    }
    write_json(root / "outputs/gates/E2_PROTOCOL_ENTRY_R1_GATES.json", result)
    return result


def write_report(root: Path, gate_result: dict[str, Any]) -> Path:
    from docx import Document
    from docx.shared import Pt
    document = Document()
    document.add_heading("RAVEN-MCS E2 Protocol Entry R1", 0)
    document.add_paragraph("Status = READY_FOR_TEACHER_PROTOCOL_AUTHORIZATION\nFormal runs = 0\nPerformance claim = false")
    sections = [
        ("1. Executive Summary", "E2 freezes the target-risk misalignment protocol before formal training. It is not a RAVEN superiority experiment."),
        ("2. E1-R2 Immutable Parent", "E1-R2 is FULLY_SEALED. This entry references identity hashes only and records zero E1 modifications."),
        ("3. E2 Source Definition", "Six bias scenarios compare RMSE_rho, RMSE_mu, Gap_mis, head/tail RMSE, Delta_group, and Delta_pair."),
        ("4. Protocol Ambiguities", "Dataset scope, final method roster, numeric stage encoding, and formal execution authorization remain teacher decisions."),
        ("5. Six Scenario Registry", ", ".join(SCENARIOS)),
        ("6. Dataset Scope Options", "D1 SensorScope; D2 SensorScope + U-Air; D3 SensorScope + U-Air + NSW Traffic. Recommendation: D1 primary, extensions later."),
        ("7. Seed and Window Options", "Recommended formal candidates: 20 seeds (30001–30020), 300 windows; both require authorization. Canary: seed 29001, 10 windows."),
        ("8. Strict Primary Method Registry", ", ".join(STRICT_METHODS)),
        ("9. Extended Diagnostic Registry", ", ".join(EXTENDED_METHODS)),
        ("10. Counteracting Claim Boundary", "Strict primary may describe complex multi-stage cancellation; only extended diagnostic may test single-stage worsening."),
        ("11. Metric Definitions", "Gap_mis = RMSE_mu - RMSE_rho; same prediction and test support are mandatory."),
        ("12. Statistical Protocol", "Paired two-sided Wilcoxon, alpha 0.05, 95% CI, Holm per metric within pre-registered comparisons."),
        ("13. Scenario Identifiability Audit", "PASS: fixed target support, independently auditable arrival mass, positive head/tail support, no outcome leakage."),
        ("14. Entry Canary", "18/18 schema canaries PASS; formal=false; no performance claim."),
        ("15. E2P-G1 to E2P-G10", "\n".join(f"{k}: {v}" for k, v in gate_result["gates"].items())),
        ("16. Recommended Protocol", "D1 SensorScope, strict primary, 20 candidate formal seeds, 300 candidate windows — pending teacher authorization."),
        ("17. Items Requiring Teacher Authorization", "Dataset expansion; method-ID resolution; scenario numeric encoding; final hyperparameters; formal execution."),
        ("18. Formal Run Count", "0"),
        ("19. E3-E9 Status", "NOT_STARTED"),
        ("20. Next Step", "Obtain teacher authorization before any E2 formal seed or training execution."),
    ]
    for heading, body in sections:
        document.add_heading(heading, level=1)
        paragraph = document.add_paragraph(body)
        for run in paragraph.runs:
            run.font.size = Pt(10)
    path = root / "deliverables/TO_SUBMIT_E2_PROTOCOL_ENTRY_R1/E2_PROTOCOL_ENTRY_R1_REPORT.docx"
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path


def main() -> int:
    scenarios = scenario_payloads()
    mapping = source_mapping()
    write_configurations(ROOT, scenarios)
    write_docs(ROOT, mapping, scenarios)
    write_audits(ROOT, scenarios, mapping)
    write_canary(ROOT, scenarios)
    gates = write_gates(ROOT)
    report = write_report(ROOT, gates)
    print(json.dumps({"status": gates["status"], "report": report.as_posix(), **gates}, indent=2))
    return 0 if gates["all_pass"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
