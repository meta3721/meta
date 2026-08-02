#!/usr/bin/env python3
"""Export the truthful E1-ENTRY-R1 partial evidence package."""
from __future__ import annotations

import subprocess
import zipfile
from pathlib import Path

import pandas as pd
from docx import Document

from raven_mcs.experiments.e1_entry import E1_SEEDS
from raven_mcs.utils.hashing import sha256_file
from raven_mcs.utils.serialization import dump_json, load_json


def _git(root: Path, *args: str) -> str:
    return subprocess.run(
        ["git", *args], cwd=root, check=True, capture_output=True, text=True,
    ).stdout


def main() -> int:
    root = Path(__file__).resolve().parents[1]
    deliverables = root / "deliverables"
    deliverables.mkdir(parents=True, exist_ok=True)
    commit = _git(root, "rev-parse", "HEAD").strip()
    status = _git(root, "status", "--porcelain")
    safety = load_json(root / "outputs/validation/e1_r1_weight_safety_report.json")
    measurement = load_json(root / "outputs/audits/e1_r1_measurement_audit.json")
    gates = load_json(root / "outputs/audits/E1_ENTRY_R1_GATE_RESULTS.json")
    pi = load_json(root / "configs/frozen/e1_pi_target_manifest.json")
    trace_manifest = load_json(
        root / "outputs/event_traces/E1_BALANCED_EVENTTRACE_MANIFEST.json",
    )
    support = pd.read_csv(root / "outputs/audits/e1_r1_group_support_by_split.csv")
    report = f"""# E1-ENTRY-R1 Report

## 1. Executive Summary

Status: **PARTIAL**. Infrastructure and semantic gates G1–G5/G7–G8 pass.
TimeAlign-Agg remains `BASELINE_UNRESOLVED`: no primary source or frozen
formula exists, and the current code is exactly FedAsync-Window. E1 remains
blocked; the five-method smoke was intentionally not represented as complete.

## 2. Formal Git Identity

- Commit: `{commit}`
- Clean at export: `{not bool(status.strip())}`
- Branch: `e1-entry-r1-semantic-seal`

## 3. Repeatable G=4 Target Groups

UTC blocks are 00–06, 06–12, 12–18 and 18–24. They replace global time-axis
quartiles.

## 4. Split Support Audit

All four groups have positive atomic and target support in train, validation
and test ({len(support)} audited split-group rows).

## 5. Client Potential Measurements

Exact `(client_id, unit_id)` lookup found
{measurement['number_potential_measurements_found']} of
{measurement['number_local_training_records']} requested records.
Missing={measurement['missing_measurement_count']};
direct target usage={measurement['direct_target_value_usage_count']}.

## 6. Calibration Residual

Delta_cal={measurement['Delta_cal']}; status
`{measurement['calibration_centering_status']}`.

## 7. Observation Propensity Information Boundary

Status `REVIEWED_WHITELIST_ONLY`; the vector is
`bias, hour_block, planned_workload_pre`. No current O outcome is used.

## 8. Frozen Client-Stratum Target Mass

pi target SHA-256: `{pi['pi_target_hash']}`; sum={pi['pi_sum']};
support violations={pi['support_violations']}.

## 9. First-Stage Weight Safety

Validation-only selection: first-stage clip
{safety['first_stage_clip_rate']}, second-stage clip
{safety['second_stage_clip_rate']}, median n_eff
{safety['median_n_eff']}. `a_max` remained 20.

## 10. TimeAlign Baseline Reconstruction

No original TimeAlign source/formula could be recovered. The implementation is
the same `n_k exp(-0.1 tau_k)` rule as FedAsync. Inventing another mechanism is
forbidden. Teacher decision required: provide the source or remove/re-freeze.

## 11. FedAsync vs TimeAlign Diagnostics

Controlled nonzero-staleness fixture: identical weights; G6 FAIL.

## 12. Lagged Empirical Variance

EMA mean and squared-deviation state are used only after window close.
Current P2 uses the pre-update state divided by n_eff plus v_floor.

## 13. Normalized Staleness

P2 receives `tau/S_max`; diagnostics preserve raw and normalized values.

## 14. Hash Identity Separation

Run manifests now separate protocol, resolved-run, data, group, client,
pi-target, trace, initial-model and environment identities.

## 15. EventTrace Evidence and Recomputed Audits

All five directories contain events.parquet and independently recomputed
audits/hashes. Trace hashes:
{chr(10).join(f"- {seed}: `{trace_manifest['traces'][str(seed)]['event_trace_hash']}`" for seed in E1_SEEDS)}

## 16. Five-Method One-Seed Smoke

Not completed. It is blocked before execution because TimeAlign has no valid
method semantics. No single-seed performance conclusion is made.

## 17. E1R1-G1 to E1R1-G10

{chr(10).join(f"- {key}: {value}" for key, value in gates['gates'].items())}

## 18. Full Test Results

Full pytest passes with three skips and one explicit expected failure that
keeps the unresolved TimeAlign requirement visible.

## 19. Remaining Issues

ISSUE-031 is unresolved. G9 cannot pass until G6 is resolved.

## 20. E1 Authorization Request

Authorization is **not requested**. E1 remains BLOCKED. Formal five-seed E1
has NOT been executed; E2–E9 have NOT started.
"""
    markdown = deliverables / "E1_ENTRY_R1_REPORT.md"
    markdown.write_text(report, encoding="utf-8")
    document = Document()
    for line in report.splitlines():
        if line.startswith("# "):
            document.add_heading(line[2:], 0)
        elif line.startswith("## "):
            document.add_heading(line[3:], 1)
        elif line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        elif line:
            document.add_paragraph(line)
    docx = deliverables / "E1_ENTRY_R1_REPORT.docx"
    document.save(docx)
    readme = deliverables / "E1_ENTRY_R1_SUBMISSION_README.txt"
    readme.write_text(
        "E1-ENTRY-R1 status: PARTIAL\n"
        "Blocker: TimeAlign-Agg has no recoverable primary definition and "
        "duplicates FedAsync-Window.\n"
        "E1 remains BLOCKED. Formal five-seed E1 and E2-E9 were not run.\n",
        encoding="utf-8",
    )
    evidence = deliverables / "RAVEN_MCS_E1_ENTRY_R1_EVIDENCE.zip"
    tracked = [
        root / path for path in _git(root, "ls-files").splitlines()
    ]
    extras = [
        root / "configs/frozen/e1_sensorscope_balanced.yaml",
        root / "configs/frozen/e1_sensorscope_clients.yaml",
        root / "configs/frozen/e1_pi_target_client_stratum.parquet",
        root / "configs/frozen/e1_pi_target_manifest.json",
        root / "configs/frozen/e1_weight_safety.yaml",
        root / "outputs/audits",
        root / "outputs/validation",
        root / "outputs/event_traces",
        root / "logs/E1_ENTRY_R1_COMMANDS.log",
        markdown, docx, readme,
    ]
    with zipfile.ZipFile(evidence, "w", zipfile.ZIP_DEFLATED) as archive:
        seen = set()
        for candidate in tracked + extras:
            paths = candidate.rglob("*") if candidate.is_dir() else [candidate]
            for path in paths:
                if not path.is_file():
                    continue
                relative = path.relative_to(root)
                if relative in seen:
                    continue
                seen.add(relative)
                archive.write(path, relative.as_posix())
    if zipfile.ZipFile(evidence).testzip() is not None:
        raise RuntimeError("evidence zip integrity failure")
    dump_json({
        "status": "PARTIAL",
        "commit": commit,
        "evidence_sha256": sha256_file(evidence),
        "report_sha256": sha256_file(docx),
        "readme_sha256": sha256_file(readme),
    }, deliverables / "E1_ENTRY_R1_DELIVERABLE_HASHES.json")
    print(evidence)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
