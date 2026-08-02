#!/usr/bin/env python3
"""Export completed E1 formal execution evidence without running experiments."""
from __future__ import annotations

import shutil
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]


def main() -> int:
    aggregate = ROOT / "outputs/aggregate/E1_balanced"
    stats = ROOT / "outputs/statistics/E1_balanced"
    gate_report = aggregate / "E1_FORMAL_GATE_REPORT.json"
    required = [aggregate / "per_seed_metrics.parquet", aggregate / "identity_audit.json",
                aggregate / "completeness_matrix.csv", stats / "wilcoxon_results.csv",
                stats / "holm_results.csv", stats / "no_harm_summary.json", gate_report]
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise RuntimeError(f"formal evidence unavailable; missing: {missing}")
    destination = ROOT / "deliverables/E1_FORMAL_EXECUTION_R1"
    if destination.exists():
        shutil.rmtree(destination)
    destination.mkdir(parents=True)
    for source in required + [ROOT / "logs/E1_FORMAL_PROGRESS.json",
                              ROOT / "configs/frozen/e1_sensorscope_balanced.yaml"]:
        if source.exists():
            shutil.copy2(source, destination / source.name)
    doc = Document()
    doc.add_heading("E1 Formal Execution R1 Evidence", level=1)
    doc.add_paragraph(f"Generated: {datetime.now(timezone.utc).isoformat()}")
    doc.add_heading("Included Evidence", level=2)
    for source in required:
        doc.add_paragraph(source.name, style="List Bullet")
    doc.save(destination / "E1_FORMAL_EXECUTION_R1_REPORT.docx")
    (destination / "README.txt").write_text(
        "Formal evidence package. See E1_FORMAL_GATE_REPORT.json for the final gate decision.\n",
        encoding="utf-8")
    zip_path = ROOT / "deliverables/RAVEN_MCS_E1_FORMAL_EXECUTION_R1_EVIDENCE.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
        for path in destination.rglob("*"):
            if path.is_file():
                archive.write(path, path.relative_to(destination).as_posix())
    print(zip_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
