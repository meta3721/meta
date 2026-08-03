#!/usr/bin/env python3
"""Export PARTIAL E1-FORMAL-EXECUTION-R1 evidence after Phase B hard-gate stop."""
from __future__ import annotations

import hashlib
import json
import shutil
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
RUN = ROOT / "outputs/runs/E1_FORMAL_fedavg_window_26001_20260802_154112_605315"
COMMIT = "c077cfb742a1f25d612228dd6881e4866ca36e04"
ALGO = "53e277c53b01695330652b8e1bc8a234909d56e5"


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> int:
    progress = json.loads((ROOT / "logs/E1_FORMAL_PROGRESS.json").read_text(encoding="utf-8"))
    manifest = json.loads((RUN / "manifest.json").read_text(encoding="utf-8"))
    metrics = json.loads((RUN / "metrics_run.json").read_text(encoding="utf-8"))
    gates = json.loads((RUN / "RUN_GATE_REPORT.json").read_text(encoding="utf-8"))
    harness = json.loads(
        (ROOT / "outputs/audits/E1_FORMAL_HARNESS_GATE_REPORT.json").read_text(encoding="utf-8")
    )

    summary = {
        "round": "E1-FORMAL-EXECUTION-R1",
        "status": "PARTIAL",
        "failure_class": "hard_gate",
        "stop_reason": "E1-RUN-G6 weight safety failed on first formal run",
        "formal_execution_commit": COMMIT,
        "authorized_algorithm_commit": ALGO,
        "protocol_parent_commit": ALGO,
        "git_clean_at_execution": True,
        "selected_baseline": "flamf_timealign_adapted",
        "harness_gates": harness.get("gates"),
        "harness_all_pass": harness.get("all_pass"),
        "formal_runs_planned": 25,
        "formal_runs_completed": 1,
        "formal_runs_passed": 0,
        "formal_runs_failed": 1,
        "failed_runs": [
            {
                "run_dir": str(RUN.relative_to(ROOT)).replace("\\", "/"),
                "method": "fedavg_window",
                "seed": 26001,
                "hard_gate_status": manifest.get("hard_gate_status"),
                "failed_gate": "E1-RUN-G6",
                "first_stage_clip_rate": metrics.get("first_stage_clip_rate"),
                "second_stage_clip_rate": metrics.get("second_stage_clip_rate"),
                "median_n_eff": metrics.get("median_n_eff"),
                "run_gates": gates.get("gates"),
            }
        ],
        "retry_count": 0,
        "aggregate_executed": False,
        "statistics_executed": False,
        "formal_gates_executed": False,
        "wilcoxon_rows": None,
        "holm_rows": None,
        "no_harm_status": "NOT_EVALUATED",
        "one_sided_95_upper_bound": None,
        "max_first_stage_clip_rate": metrics.get("first_stage_clip_rate"),
        "max_second_stage_clip_rate": metrics.get("second_stage_clip_rate"),
        "min_median_n_eff": metrics.get("median_n_eff"),
        "q_nonattempt_leakage_total": metrics.get("q_nonattempt_leakage_count"),
        "failed_attempt_omission_total": metrics.get("q_failed_attempt_omission_count"),
        "unsupported_arrival_contribution_sum": metrics.get(
            "unsupported_arrival_contribution_sum"
        ),
        "solver_failure_total": metrics.get("solver_failure_count"),
        "E1_final_status": "PARTIAL_STOPPED_ON_RUN_HARD_GATE",
        "E2_E9_status": "NOT_STARTED",
        "progress": progress,
        "generated_utc": datetime.now(timezone.utc).isoformat(),
    }
    summary_path = ROOT / "outputs/audits/E1_FORMAL_EXECUTION_R1_PARTIAL_SUMMARY.json"
    summary_path.write_text(json.dumps(summary, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    cmd_log = ROOT / "logs/E1_FORMAL_EXECUTION_R1_EXACT_COMMANDS.jsonl"
    entry = {
        "stage": "formal_25_run_fail_fast",
        "command": (
            f"{ROOT / '.venv' / 'Scripts' / 'python.exe'} scripts/run_e1_formal.py "
            "--config configs/frozen/e1_sensorscope_balanced.yaml "
            "--methods fedavg_window fedasync_window flamf_timealign_adapted "
            "twostage_hajek raven --seeds 26001 26002 26003 26004 26005 "
            "--windows 100 --device cpu --output-root outputs/runs --fail-fast"
        ),
        "cwd": str(ROOT),
        "start_time": progress.get("start_time"),
        "end_time": progress.get("last_update_time"),
        "exit_code": 1,
        "stdout_log": None,
        "stderr_log": None,
        "output_paths": [
            "logs/E1_FORMAL_PROGRESS.json",
            str(RUN.relative_to(ROOT)).replace("\\", "/"),
            str(summary_path.relative_to(ROOT)).replace("\\", "/"),
        ],
        "output_hashes": [],
        "note": (
            "fail-fast after E1-RUN-G6 on fedavg_window/26001; "
            "remaining 24 runs not started"
        ),
    }
    with cmd_log.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(entry, ensure_ascii=False) + "\n")

    dest = ROOT / "deliverables/E1_FORMAL_EXECUTION_R1_PARTIAL"
    if dest.exists():
        shutil.rmtree(dest)
    dest.mkdir(parents=True)

    for src in [
        ROOT / "logs/E1_FORMAL_PROGRESS.json",
        ROOT / "outputs/audits/E1_FORMAL_HARNESS_GATE_REPORT.json",
        summary_path,
        cmd_log,
        ROOT / "configs/frozen/e1_sensorscope_balanced.yaml",
        ROOT / "configs/frozen/e1_weight_safety.yaml",
        ROOT / "configs/frozen/e1_selected_baseline.yaml",
        ROOT / "configs/frozen/FROZEN_CONFIG_MANIFEST.json",
        RUN / "manifest.json",
        RUN / "metrics_run.json",
        RUN / "RUN_GATE_REPORT.json",
        RUN / "RUN_GATE_REPORT.md",
        RUN / "resolved_config.yaml",
        RUN / "event_trace_ref.json",
        RUN / "system_metrics.json",
    ]:
        if src.exists():
            shutil.copy2(src, dest / src.name)

    run_pack = dest / "failed_run_E1_FORMAL_fedavg_window_26001"
    run_pack.mkdir()
    for name in [
        "manifest.json",
        "metrics_run.json",
        "metrics_window.parquet",
        "RUN_GATE_REPORT.json",
        "RUN_GATE_REPORT.md",
        "resolved_config.yaml",
        "event_trace_ref.json",
        "predictions_test.parquet",
        "arrival_weights_test.parquet",
        "q_attempt_diagnostics.parquet",
        "arrival_support_diagnostics.parquet",
        "solver_diagnostics.parquet",
        "propensity_diagnostics.parquet",
        "opportunity_ema_diagnostics.parquet",
        "method_diagnostics.parquet",
        "calibration_summary.json",
        "feature_encoding_manifest.json",
        "system_metrics.json",
        "stdout.log",
        "stderr.log",
    ]:
        path = RUN / name
        if path.exists():
            shutil.copy2(path, run_pack / name)
    ckpt = RUN / "checkpoints/final.pt"
    if ckpt.exists():
        (run_pack / "checkpoints").mkdir(exist_ok=True)
        shutil.copy2(ckpt, run_pack / "checkpoints/final.pt")

    doc = Document()
    doc.add_heading("E1-FORMAL-EXECUTION-R1 Report (PARTIAL)", level=1)
    doc.add_heading("1. Executive Summary", level=2)
    doc.add_paragraph(
        "Phase A harness sealed and HARNESS-G1..G10 all PASS. Phase B started the "
        "formal 5x5x100 matrix under fail-fast and stopped after the first run: "
        f"fedavg_window / seed 26001 failed E1-RUN-G6 "
        f"(first_stage_clip_rate={metrics['first_stage_clip_rate']:.6f} > 0.05). "
        "Remaining 24 runs were not started. No parameter retuning and no seed "
        "replacement were performed."
    )
    doc.add_heading("2. Formal Execution Commit", level=2)
    doc.add_paragraph(f"formal_execution_commit = {COMMIT}")
    doc.add_paragraph(f"authorized_algorithm_commit = {ALGO}")
    doc.add_paragraph(f"protocol_parent_commit = {ALGO}")
    doc.add_heading("3. Frozen Protocol", level=2)
    for line in [
        "dataset=SensorScope, scenario=balanced, clients=8, G=4, S_max=5",
        "local_steps=2, num_windows=100, no_harm_threshold=0.03",
        "selected_baseline=flamf_timealign_adapted",
        "seeds=26001..26005; methods=fedavg/fedasync/flamf/twostage/raven",
    ]:
        doc.add_paragraph(line, style="List Bullet")
    doc.add_heading("4. Harness Gate Results", level=2)
    for key, value in harness.get("gates", {}).items():
        doc.add_paragraph(f"{key}: {value}", style="List Bullet")
    doc.add_heading("5. 25-Run Completion Matrix", level=2)
    doc.add_paragraph("Completed=1/25, Passed=0, Failed=1, Not started=24")
    doc.add_heading("6. Per-Run Gate Summary", level=2)
    for gate in gates["gates"]:
        doc.add_paragraph(
            f"{gate['gate']}: {gate['status']} — {gate['detail']}",
            style="List Bullet",
        )
    doc.add_heading("7. Per-Seed RMSE_mu", level=2)
    doc.add_paragraph(
        "Only failed fedavg_window@26001 available: "
        f"RMSE_mu={metrics['RMSE_mu']:.6f}"
    )
    doc.add_heading("8. RMSE_rho and Gap_mis", level=2)
    doc.add_paragraph(
        f"RMSE_rho={metrics['RMSE_rho']:.6f}; Gap_mis={metrics['Gap_mis']:.6f}"
    )
    doc.add_heading("9. Head/Tail Results", level=2)
    doc.add_paragraph(
        f"Head_RMSE={metrics['Head_RMSE']:.6f}; Tail_RMSE={metrics['Tail_RMSE']:.6f}"
    )
    doc.add_heading("10. No-Harm Paired Degradation", level=2)
    doc.add_paragraph("NOT_EVALUATED (matrix incomplete)")
    doc.add_heading("11. One-Sided 95% Upper Bound", level=2)
    doc.add_paragraph("NOT_EVALUATED")
    doc.add_heading("12. Wilcoxon Results", level=2)
    doc.add_paragraph("NOT_PRODUCED (aggregate blocked by incomplete formal set)")
    doc.add_heading("13. Holm-Corrected Results", level=2)
    doc.add_paragraph("NOT_PRODUCED")
    doc.add_heading("14. Clip and ESS Safety", level=2)
    doc.add_paragraph(
        f"first_stage_clip_rate={metrics['first_stage_clip_rate']:.6f} "
        f"(FAIL threshold 0.05); "
        f"second_stage_clip_rate={metrics['second_stage_clip_rate']:.6f}; "
        f"median_n_eff={metrics['median_n_eff']:.6f}"
    )
    doc.add_heading("15. Stage-2 Attempt Audit", level=2)
    doc.add_paragraph(
        f"q_nonattempt_leakage_count={metrics.get('q_nonattempt_leakage_count')}; "
        f"q_failed_attempt_omission_count="
        f"{metrics.get('q_failed_attempt_omission_count')}"
    )
    doc.add_heading("16. Arrival Support Audit", level=2)
    doc.add_paragraph(
        "unsupported_arrival_contribution_count="
        f"{metrics.get('unsupported_arrival_contribution_count')}; "
        f"sum={metrics.get('unsupported_arrival_contribution_sum')}"
    )
    doc.add_heading("17. Solver Reliability", level=2)
    doc.add_paragraph(
        f"solver_failure_count={metrics.get('solver_failure_count')}; "
        f"fallback_count={metrics.get('fallback_count')}"
    )
    doc.add_heading("18. Runtime and Communication", level=2)
    doc.add_paragraph(
        f"total_runtime={metrics.get('total_runtime'):.3f}s; "
        f"total_communication={metrics.get('total_communication')}"
    )
    doc.add_heading("19. Failed or Retried Runs", level=2)
    doc.add_paragraph(
        "Failed: E1_FORMAL_fedavg_window_26001_20260802_154112_605315 "
        "(hard_gate, retained). Retries=0."
    )
    doc.add_heading("20. E1-G1 to E1-G6", level=2)
    doc.add_paragraph(
        "Formal aggregate gates not executed because 25/25 PASS runs were not achieved."
    )
    doc.add_heading("21. E1 Final Decision", level=2)
    doc.add_paragraph(
        "PARTIAL — Phase B stopped on first-run weight-safety hard gate; "
        "formal performance claim not authorized."
    )
    doc.add_heading("22. E2-E9 Status", level=2)
    doc.add_paragraph("NOT_STARTED")
    doc.add_heading("23. Next Round Recommendation", level=2)
    doc.add_paragraph(
        "Teacher decision required: formal 100-window first-stage clip on "
        "fedavg_window@26001 exceeds the frozen 5% gate under the "
        "validation-selected weight-safety config. Do not retune in this round."
    )

    report_docx = ROOT / "deliverables/E1_FORMAL_R1_REPORT.docx"
    doc.save(report_docx)
    doc.save(dest / "E1_FORMAL_R1_REPORT.docx")

    readme = f"""RAVEN-MCS E1-FORMAL-EXECUTION-R1 submission (PARTIAL)
====================================================
status=PARTIAL
formal_execution_commit={COMMIT}
authorized_algorithm_commit={ALGO}
protocol_parent_commit={ALGO}
git_clean_at_execution=True
selected_baseline=flamf_timealign_adapted
local_steps=2
num_windows=100
formal_runs_completed=1
formal_runs_passed=0
formal_runs_failed=1
failed_run=outputs/runs/E1_FORMAL_fedavg_window_26001_20260802_154112_605315
failed_gate=E1-RUN-G6
first_stage_clip_rate={metrics['first_stage_clip_rate']}
second_stage_clip_rate={metrics['second_stage_clip_rate']}
median_n_eff={metrics['median_n_eff']}
retries=0
aggregate=NOT_RUN
statistics=NOT_RUN
E1=PARTIAL_STOPPED_ON_RUN_HARD_GATE
E2-E9=NOT_STARTED
notes=Phase A HARNESS-G1..G10 PASS. Phase B fail-fast retained the FAIL directory and did not start the remaining 24 runs. No retuning.
"""
    readme_path = ROOT / "deliverables/E1_FORMAL_R1_SUBMISSION_README.txt"
    readme_path.write_text(readme, encoding="utf-8")
    (dest / "E1_FORMAL_R1_SUBMISSION_README.txt").write_text(readme, encoding="utf-8")

    zip_path = ROOT / "deliverables/RAVEN_MCS_E1_FORMAL_R1_EVIDENCE.zip"
    if zip_path.exists():
        zip_path.unlink()
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
        for path in dest.rglob("*"):
            if path.is_file():
                archive.write(path, path.relative_to(dest).as_posix())

    hashes = {
        str(path.relative_to(ROOT)).replace("\\", "/"): _sha(path)
        for path in [zip_path, report_docx, readme_path, summary_path]
    }
    (ROOT / "deliverables/E1_FORMAL_R1_DELIVERABLE_HASHES.json").write_text(
        json.dumps(hashes, indent=2) + "\n", encoding="utf-8"
    )
    export_entry = {
        "stage": "partial_evidence_export",
        "command": f"{ROOT / '.venv' / 'Scripts' / 'python.exe'} scripts/export_e1_formal_r1_partial_evidence.py",
        "cwd": str(ROOT),
        "start_time": datetime.now(timezone.utc).isoformat(),
        "end_time": datetime.now(timezone.utc).isoformat(),
        "exit_code": 0,
        "stdout_log": None,
        "stderr_log": None,
        "output_paths": list(hashes),
        "output_hashes": list(hashes.values()),
    }
    with cmd_log.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(export_entry, ensure_ascii=False) + "\n")
    print(json.dumps({"zip": str(zip_path), "hashes": hashes}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
